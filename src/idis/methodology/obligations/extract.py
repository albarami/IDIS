"""Workbook and template extraction into a private ``derived/`` directory."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from idis.methodology.category_mapping import load_commercial_dd_question_category_map
from idis.methodology.ids import normalize_text
from idis.methodology.importers.cdd_json import import_cdd_template
from idis.methodology.importers.fdd_excel import (
    PRODUCTION_COLUMNS,
    FDDImporterError,
    validated_production_fdd_workbook_projection,
)
from idis.methodology.obligations.errors import Gl02Error
from idis.methodology.obligations.extract_cells import (
    _cell_key,
    _conditional_count,
    _feature_present,
    _input_bytes,
    _map_ids,
    _membership,
    _parent_key,
    _parity,
    _read_bytes,
    _source_cell,
    _text,
)
from idis.methodology.obligations.jsonl import sha256_bytes, sha256_text, write_json, write_jsonl
from idis.methodology.obligations.models_records import (
    CellRef,
    CommercialContext,
    EdgeWhitespaceCell,
    ExtractionReport,
    FeatureCensus,
    InputFingerprint,
    MergedInstruction,
    ParentRecord,
    PolicyViewRecord,
    ReportMappingRef,
    SheetCount,
    SourceCell,
)
from idis.methodology.obligations.policy_view import build_policy_view
from idis.methodology.obligations.vocab import (
    REGISTER_SCHEMA,
    REGISTER_VERSION,
    SHEET_CODES,
    SHEET_NAMES,
    Code,
    ParentKind,
)

_COMMERCIAL_ID = re.compile(r"^mq_commercial_dd_(.+)_0001$")
_TEMPLATE_PACKAGE = "idis.methodology"
_TEMPLATE_NAME = "templates/commercial_dd_v1.json"
_MAP_PACKAGE = "idis.methodology.mappings"
_MAP_NAME = "commercial_dd_v1_question_categories.json"


def extract_register(
    workbook_path: Path | str,
    private_root: Path | str,
    *,
    template_path: Path | str | None = None,
    category_map_path: Path | str | None = None,
    policy_path: Path | str | None = None,
) -> ExtractionReport:
    """Extract derived files. The same inputs yield byte-identical files."""
    source = Path(workbook_path)
    root = Path(private_root)
    before = _read_bytes(source)
    workbook_fp = InputFingerprint(sha256=sha256_bytes(before), byte_count=len(before))
    try:
        projection = validated_production_fdd_workbook_projection(before)
    except FDDImporterError as exc:
        raise Gl02Error(Code.GL02_SOURCE_COVERAGE.value) from exc
    template_bytes = _input_bytes(template_path, _TEMPLATE_PACKAGE, _TEMPLATE_NAME)
    template_fp = InputFingerprint(
        sha256=sha256_bytes(template_bytes), byte_count=len(template_bytes)
    )
    map_bytes = _input_bytes(category_map_path, _MAP_PACKAGE, _MAP_NAME)
    map_fp = InputFingerprint(sha256=sha256_bytes(map_bytes), byte_count=len(map_bytes))
    registry = import_cdd_template(template_bytes)
    category_map = load_commercial_dd_question_category_map()
    loaded = load_workbook(BytesIO(before), data_only=False, read_only=False)
    try:
        cells, census, merges = _census(loaded)
        parents, parity_ids, instructions = _parents(
            loaded,
            projection_rows=projection.rows,
            workbook_hash=workbook_fp.sha256,
            questions=registry.current_version.questions,
            category_rows=category_map.questions,
            version_id=category_map.methodology_version_id,
            source_hash=category_map.methodology_source_sha256,
            merges=merges,
        )
    finally:
        loaded.close()
    after = _read_bytes(source)
    if after != before:
        raise Gl02Error(Code.GL02_SOURCE_CHANGED_DURING_READ.value)
    findings = list(parity_ids and (Code.GL02_IMPORTER_PARITY_FAILED.value,) or ())
    if _feature_present(census):
        findings.append(Code.GL02_SOURCE_FEATURE_PRESENT.value)
    if census.other_role_stored_cells:
        findings.append(Code.GL02_SOURCE_OTHER_CELL.value)
    if category_map_path is not None and _map_ids(map_bytes) != tuple(
        row.methodology_question_id for row in category_map.questions
    ):
        findings.append(Code.GL02_CATEGORY_MAP_MISMATCH.value)
    policy_fp = None
    policy_records: tuple[PolicyViewRecord, ...] = ()
    anchor_count = 0
    source_count = 0
    non_source = 0
    if policy_path is not None:
        built = build_policy_view(Path(policy_path), tuple(parents))
        policy_fp = built.fingerprint
        policy_records = built.records
        anchor_count = built.anchor_count
        source_count = built.source_binding_count
        non_source = built.non_source_input_spec_count
        findings.extend(finding.code for finding in built.findings)
    report = ExtractionReport(
        register_schema=REGISTER_SCHEMA,
        register_version=REGISTER_VERSION,
        workbook=workbook_fp,
        template=template_fp,
        category_map=map_fp,
        policy=policy_fp,
        sheet_names=SHEET_NAMES,
        data_row_counts=tuple(
            SheetCount(
                sheet=name, count=sum(1 for row in projection.rows if row.source_sheet == name)
            )
            for name in SHEET_NAMES
        ),
        merged_range_counts=tuple(
            SheetCount(sheet=name, count=sum(1 for item in merges if item[0] == name))
            for name in SHEET_NAMES
        ),
        feature_census=census,
        edge_whitespace_cells=tuple(
            EdgeWhitespaceCell(cell_id=cell.cell_id, role=cell.role)
            for cell in cells
            if cell.edge_whitespace
        ),
        parity_failure_ids=tuple(sorted(parity_ids)),
        anchor_count=anchor_count,
        source_binding_count=source_count,
        non_source_input_spec_count=non_source,
        finding_codes=tuple(sorted(set(findings))),
    )
    derived = root / "derived"
    derived.mkdir(parents=True, exist_ok=True)
    write_jsonl(derived / "source_cells.jsonl", cells, key=_cell_key)
    write_jsonl(derived / "parents.jsonl", parents, key=_parent_key)
    write_jsonl(
        derived / "merged_instructions.jsonl", instructions, key=lambda item: item.instruction_id
    )
    write_jsonl(derived / "policy_view.jsonl", policy_records, key=lambda item: item.parent_id)
    write_json(derived / "extraction_report.private.json", report.model_dump(mode="json"))
    return report


def _census(
    workbook: Workbook,
) -> tuple[list[SourceCell], FeatureCensus, list[tuple[str, object]]]:
    cells: list[SourceCell] = []
    comments = hyperlinks = hidden_rows = hidden_columns = 0
    validations = conditional = images = formulas = non_strings = other_stored = 0
    merges: list[tuple[str, object]] = []
    non_visible = sum(1 for sheet in workbook.worksheets if sheet.sheet_state != "visible")
    defined_names = len(tuple(workbook.defined_names.keys()))
    for ordinal, sheet_name in enumerate(SHEET_NAMES, start=1):
        sheet = workbook[sheet_name]
        hidden_rows += sum(1 for dim in sheet.row_dimensions.values() if dim.hidden)
        hidden_columns += sum(1 for dim in sheet.column_dimensions.values() if dim.hidden)
        validations += len(tuple(sheet.data_validations.dataValidation))
        conditional += _conditional_count(sheet)
        images += len(tuple(getattr(sheet, "_images", ()) or ()))
        membership = _membership(sheet, sheet_name)
        merges.extend((sheet_name, merged) for merged in sheet.merged_cells.ranges)
        positions = set(membership)
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    comments += 1
                if cell.hyperlink is not None:
                    hyperlinks += 1
                if cell.value is not None:
                    positions.add((cell.row, cell.column))
        for row_index, column_index in sorted(positions):
            cell = sheet.cell(row=row_index, column=column_index)
            record, formula, non_string, other = _source_cell(
                sheet_name, ordinal, cell, membership.get((row_index, column_index))
            )
            formulas += int(formula)
            non_strings += int(non_string)
            other_stored += int(other)
            cells.append(record)
    census = FeatureCensus(
        comments=comments,
        hyperlinks=hyperlinks,
        hidden_rows=hidden_rows,
        hidden_columns=hidden_columns,
        data_validations=validations,
        conditional_formats=conditional,
        images=images,
        defined_names=defined_names,
        formula_cells=formulas,
        non_string_stored_values=non_strings,
        non_visible_sheets=non_visible,
        other_role_stored_cells=other_stored,
    )
    return cells, census, merges


def _parents(
    workbook: Workbook,
    *,
    projection_rows: tuple[object, ...],
    workbook_hash: str,
    questions: tuple[object, ...],
    category_rows: tuple[object, ...],
    version_id: str,
    source_hash: str,
    merges: list[tuple[str, object]],
) -> tuple[list[ParentRecord], set[str], list[MergedInstruction]]:
    instruction_by_row, instructions = _merged_instructions(projection_rows, merges)
    parents: list[ParentRecord] = []
    parity: set[str] = set()
    for row in projection_rows:
        sheet = workbook[row.source_sheet]  # type: ignore[attr-defined]
        built, failed = _financial_parent(sheet, row, workbook_hash, instruction_by_row)
        parents.append(built)
        if failed:
            parity.add(row.methodology_question_id)  # type: ignore[attr-defined]
    categories = {item.methodology_question_id: item for item in category_rows}  # type: ignore[attr-defined]
    for ordinal, question in enumerate(questions, start=1):
        parents.append(_commercial_parent(question, categories, ordinal, version_id, source_hash))
    return parents, parity, instructions


def _financial_parent(
    sheet: Worksheet,
    row: object,
    workbook_hash: str,
    instruction_by_row: dict[tuple[str, int], tuple[str, ...]],
) -> tuple[ParentRecord, bool]:
    sheet_name = row.source_sheet  # type: ignore[attr-defined]
    source_row = row.source_row  # type: ignore[attr-defined]
    columns = PRODUCTION_COLUMNS[sheet_name]
    resolved = {
        name: _resolved_role(sheet, sheet_name, source_row, column)
        for name, column in columns.items()
    }
    failed = False
    for name, imported in (
        ("term", row.term),  # type: ignore[attr-defined]
        ("nature", row.nature),  # type: ignore[attr-defined]
        ("line_item", row.line_item),  # type: ignore[attr-defined]
        ("question", row.question_text),  # type: ignore[attr-defined]
    ):
        raw, _ref, anchor = resolved[name]
        if not _parity(raw, imported):
            failed = True
        if name == "question" and anchor != row.source_anchor_cell:  # type: ignore[attr-defined]
            failed = True
    question_raw = resolved["question"][0] or ""
    edge_roles = tuple(
        name
        for name, (raw, _ref, _anchor) in resolved.items()
        if raw is not None and raw != raw.strip()
    )
    parent = ParentRecord(
        parent_id=row.methodology_question_id,  # type: ignore[attr-defined]
        public_coordinate=f"FDD-{SHEET_CODES[sheet_name]}-{source_row:03d}",
        parent_kind=ParentKind.FINANCIAL,
        methodology_version_id=f"financial_dd:{workbook_hash}",
        source_hash=workbook_hash,
        sheet=sheet_name,
        source_row=source_row,
        row_ordinal=row.row_ordinal,  # type: ignore[attr-defined]
        sheet_ordinal=row.sheet_ordinal,  # type: ignore[attr-defined]
        term_raw=resolved["term"][0],
        nature_raw=resolved["nature"][0],
        line_item_raw=resolved["line_item"][0],
        question_raw=question_raw,
        term_cell=resolved["term"][1],
        nature_cell=resolved["nature"][1],
        line_item_cell=resolved["line_item"][1],
        question_cell=resolved["question"][1],
        question_raw_sha256=sha256_text(question_raw),
        question_normalized=normalize_text(question_raw),
        question_normalized_sha256=sha256_text(normalize_text(question_raw)),
        merged_instruction_ids=instruction_by_row.get((sheet_name, source_row), ()),
        edge_whitespace_roles=edge_roles,
        commercial_context=None,
    )
    return parent, failed


def _commercial_parent(
    question: object,
    categories: dict[str, object],
    ordinal: int,
    version_id: str,
    source_hash: str,
) -> ParentRecord:
    question_id = question.methodology_question_id  # type: ignore[attr-defined]
    match = _COMMERCIAL_ID.fullmatch(question_id)
    if match is None:
        raise Gl02Error(Code.GL02_PARENT_COVERAGE.value)
    question_raw = question.question_text  # type: ignore[attr-defined]
    entry = categories.get(question_id)
    governed = () if entry is None else tuple(category.value for category in entry.categories)  # type: ignore[attr-defined]
    mapping = question.report_mapping  # type: ignore[attr-defined]
    return ParentRecord(
        parent_id=question_id,
        public_coordinate=f"CDD-{match.group(1).upper()}",
        parent_kind=ParentKind.COMMERCIAL,
        methodology_version_id=version_id,
        source_hash=source_hash,
        sheet=None,
        source_row=None,
        row_ordinal=ordinal,
        sheet_ordinal=None,
        term_raw=question.term,  # type: ignore[attr-defined]
        nature_raw=question.nature,  # type: ignore[attr-defined]
        line_item_raw=question.line_item,  # type: ignore[attr-defined]
        question_raw=question_raw,
        term_cell=None,
        nature_cell=None,
        line_item_cell=None,
        question_cell=None,
        question_raw_sha256=sha256_text(question_raw),
        question_normalized=normalize_text(question_raw),
        question_normalized_sha256=sha256_text(normalize_text(question_raw)),
        merged_instruction_ids=(),
        edge_whitespace_roles=(),
        commercial_context=CommercialContext(
            report_mapping=ReportMappingRef(
                report_section=mapping.report_section,
                report_subsection=mapping.report_subsection,
            ),
            template_role=tuple(agent.role for agent in question.assigned_agents),  # type: ignore[attr-defined]
            target_document_categories=tuple(question.target_document_categories),  # type: ignore[attr-defined]
            governed_categories=governed,
        ),
    )


def _merged_instructions(
    projection_rows: tuple[object, ...],
    merges: list[tuple[str, object]],
) -> tuple[dict[tuple[str, int], tuple[str, ...]], list[MergedInstruction]]:
    by_position = {
        (row.source_sheet, row.source_row): row.methodology_question_id  # type: ignore[attr-defined]
        for row in projection_rows
    }
    grouped: dict[tuple[str, int], list[str]] = {}
    records: list[MergedInstruction] = []
    for sheet_name, merged in merges:
        question_column = PRODUCTION_COLUMNS[sheet_name]["question"]
        if question_column is None or merged.max_row <= merged.min_row:  # type: ignore[attr-defined]
            continue
        if not (merged.min_col <= question_column <= merged.max_col):  # type: ignore[attr-defined]
            continue
        anchor = f"{get_column_letter(merged.min_col)}{merged.min_row}"  # type: ignore[attr-defined]
        instruction_id = f"GL02-SI-{SHEET_CODES[sheet_name]}-{anchor}"
        covered_rows = tuple(
            row
            for row in range(merged.min_row, merged.max_row + 1)  # type: ignore[attr-defined]
            if (sheet_name, row) in by_position
        )
        covered_ids = tuple(by_position[(sheet_name, row)] for row in covered_rows)
        records.append(
            MergedInstruction(
                instruction_id=instruction_id,
                sheet=sheet_name,
                sheet_code=SHEET_CODES[sheet_name],
                anchor_cell=anchor,
                merge_range=merged.coord,  # type: ignore[attr-defined]
                covered_rows=covered_rows,
                covered_parent_ids=covered_ids,
            )
        )
        for row in covered_rows:
            grouped.setdefault((sheet_name, row), []).append(instruction_id)
    return {key: tuple(sorted(value)) for key, value in grouped.items()}, records


def _resolved_role(
    sheet: Worksheet,
    sheet_name: str,
    row_index: int,
    column: int | None,
) -> tuple[str | None, CellRef | None, str]:
    if column is None:
        return None, None, ""
    cell = sheet.cell(row=row_index, column=column)
    member = _membership(sheet, sheet_name).get((row_index, column))
    letter = get_column_letter(column)
    cell_id = f"{sheet_name}!{letter}{row_index}"
    if member is None:
        raw = _text(cell.value)
        return (
            raw,
            CellRef(cell_id=cell_id, merge_anchor=None, merge_range=None),
            f"{letter}{row_index}",
        )
    anchor = sheet.cell(row=member.anchor_row, column=member.anchor_column)
    raw = _text(anchor.value)
    ref = CellRef(cell_id=cell_id, merge_anchor=member.anchor_a1, merge_range=member.coord)
    return raw, ref, member.anchor_a1
