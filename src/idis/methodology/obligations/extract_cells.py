"""Cell census helpers for workbook extraction."""

from __future__ import annotations

import json
from dataclasses import dataclass
from importlib.resources import files
from pathlib import Path

from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from idis.methodology.importers.fdd_excel import (
    FIRST_DATA_ROW,
    HEADER_ROW,
    PRODUCTION_COLUMNS,
)
from idis.methodology.obligations.errors import Gl02Error
from idis.methodology.obligations.jsonl import sha256_text
from idis.methodology.obligations.models_records import FeatureCensus, ParentRecord, SourceCell
from idis.methodology.obligations.vocab import CellRole, Code, ParentKind


@dataclass(frozen=True, slots=True)
class _MergeMember:
    coord: str
    anchor_a1: str
    anchor_id: str
    anchor_row: int
    anchor_column: int
    is_anchor: bool


def _source_cell(
    sheet_name: str,
    ordinal: int,
    cell: Cell | MergedCell,
    member: _MergeMember | None,
) -> tuple[SourceCell, bool, bool, bool]:
    row_index = cell.row
    column_index = cell.column
    if not isinstance(row_index, int) or not isinstance(column_index, int):
        raise Gl02Error(Code.GL02_IO.value)
    role = _role(sheet_name, row_index, column_index)
    formula = cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))
    non_string = cell.value is not None and not isinstance(cell.value, str)
    other = cell.value is not None and role is CellRole.OTHER
    raw: str | None
    resolved_from: str | None
    merge_range: str | None
    merge_anchor: str | None
    is_anchor: bool
    if member is not None and not member.is_anchor:
        raw = None
        resolved_from = member.anchor_id
        merge_range = member.coord
        merge_anchor = member.anchor_a1
        is_anchor = False
    else:
        raw = _text(cell.value) if cell.value is not None else None
        resolved_from = None
        merge_range = member.coord if member else None
        merge_anchor = member.anchor_a1 if member else None
        is_anchor = member.is_anchor if member else False
    record = SourceCell(
        cell_id=f"{sheet_name}!{get_column_letter(column_index)}{row_index}",
        sheet=sheet_name,
        sheet_ordinal=ordinal,
        row=row_index,
        column=column_index,
        column_letter=get_column_letter(column_index),
        role=role,
        raw_text=raw,
        raw_sha256=None if raw is None else sha256_text(raw),
        merge_range=merge_range,
        merge_anchor=merge_anchor,
        is_merge_anchor=is_anchor,
        resolved_from=resolved_from,
        edge_whitespace=raw is not None and raw != raw.strip(),
        newline_count=0 if raw is None else raw.count("\n") + raw.count("\r"),
    )
    return record, formula and cell.value is not None, non_string, other


def _membership(sheet: Worksheet, sheet_name: str) -> dict[tuple[int, int], _MergeMember]:
    members: dict[tuple[int, int], _MergeMember] = {}
    for merged in sheet.merged_cells.ranges:
        anchor_letter = get_column_letter(merged.min_col)
        anchor_a1 = f"{anchor_letter}{merged.min_row}"
        anchor_id = f"{sheet_name}!{anchor_a1}"
        for row_index in range(merged.min_row, merged.max_row + 1):
            for column_index in range(merged.min_col, merged.max_col + 1):
                members[(row_index, column_index)] = _MergeMember(
                    coord=str(merged.coord),
                    anchor_a1=anchor_a1,
                    anchor_id=anchor_id,
                    anchor_row=merged.min_row,
                    anchor_column=merged.min_col,
                    is_anchor=row_index == merged.min_row and column_index == merged.min_col,
                )
    return members


def _role(sheet_name: str, row_index: int, column_index: int) -> CellRole:
    if row_index == HEADER_ROW:
        return CellRole.HEADER
    if row_index >= FIRST_DATA_ROW:
        for name, column in PRODUCTION_COLUMNS[sheet_name].items():
            if column == column_index:
                return CellRole(name)
    return CellRole.OTHER


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return format(value, ".15g")
    return type(value).__name__


def _parity(raw: str | None, imported: str | None) -> bool:
    if raw is None:
        return imported is None or imported == ""
    if imported is None:
        return raw.strip() == ""
    return raw.strip() == imported


def _feature_present(census: FeatureCensus) -> bool:
    return any(
        (
            census.comments,
            census.hyperlinks,
            census.hidden_rows,
            census.hidden_columns,
            census.data_validations,
            census.conditional_formats,
            census.images,
            census.defined_names,
            census.formula_cells,
            census.non_string_stored_values,
            census.non_visible_sheets,
        )
    )


def _conditional_count(sheet: Worksheet) -> int:
    rules = getattr(sheet.conditional_formatting, "_cf_rules", {})
    if not isinstance(rules, dict):
        return 0
    return sum(len(tuple(items)) for items in rules.values())


def _map_ids(raw: bytes) -> tuple[str, ...] | None:

    try:
        payload = json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return None
    questions = payload.get("questions") if isinstance(payload, dict) else None
    if not isinstance(questions, list):
        return None
    ids: list[str] = []
    for item in questions:
        if not isinstance(item, dict) or not isinstance(item.get("methodology_question_id"), str):
            return None
        ids.append(item["methodology_question_id"])
    return tuple(ids)


def _input_bytes(path: Path | str | None, package: str, name: str) -> bytes:
    if path is None:
        return files(package).joinpath(name).read_bytes()
    return _read_bytes(Path(path))


def _read_bytes(path: Path) -> bytes:
    try:
        return path.read_bytes()
    except OSError as exc:
        raise Gl02Error(Code.GL02_IO.value) from exc


def _cell_key(record: SourceCell) -> tuple[int, int, int]:
    return (record.sheet_ordinal, record.row, record.column)


def _parent_key(record: ParentRecord) -> tuple[int, int, int, str]:
    kind = 0 if record.parent_kind is ParentKind.FINANCIAL else 1
    return (kind, record.sheet_ordinal or 0, record.row_ordinal, record.parent_id)
