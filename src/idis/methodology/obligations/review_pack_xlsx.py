"""Shard workbooks for the private review pack.

Every string is stored with an explicit string data type so a leading
``=``, ``+``, ``-`` or ``@`` cannot become a formula.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from idis.methodology.obligations.models_authored import ChildObligation
from idis.methodology.obligations.register_io import Register
from idis.methodology.obligations.review_common import (
    disposed_text,
    ordered_parents,
    source_coordinate,
)
from idis.methodology.obligations.runs import field_text
from idis.methodology.obligations.vocab import Shard

_WRAP = Alignment(wrap_text=True, vertical="top")
_TEXT_WIDTH = 36
_ID_WIDTH = 22
_CELL_MAX_CHARS = 32767
_CELL_CUT_SUFFIX = "[GL02-CUT]"
_C0_CONTROL_LIMIT = 32


def write_shard_workbook(path: Path, register: Register, shard: Shard) -> None:
    """Write one deterministic workbook for a shard."""
    workbook = Workbook()
    default = workbook.active
    parents = workbook.create_sheet("Parents")
    obligations = workbook.create_sheet("Obligations")
    ambiguities = workbook.create_sheet("Ambiguities")
    cross_row = workbook.create_sheet("CrossRow")
    dispositions = workbook.create_sheet("Dispositions")
    if isinstance(default, Worksheet):
        workbook.remove(default)
    _parents(parents, register, shard)
    _obligations(obligations, register, shard)
    _ambiguities(ambiguities, register, shard)
    _cross_rows(cross_row, register, shard)
    _dispositions(dispositions, register, shard)
    for sheet in workbook.worksheets:
        _finish(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def _parents(sheet: Worksheet, register: Register, shard: Shard) -> None:
    _header(
        sheet,
        (
            "public_coordinate",
            "sheet",
            "row_ordinal",
            "parent_id",
            "question",
            "line_item",
            "term",
            "nature",
        ),
    )
    for parent in ordered_parents(register, (shard,)):
        _row(
            sheet,
            [
                parent.public_coordinate,
                parent.sheet or "",
                parent.row_ordinal,
                parent.parent_id,
                field_text(parent, "question") or "",
                field_text(parent, "line_item") or "",
                field_text(parent, "term") or "",
                field_text(parent, "nature") or "",
            ],
        )


def _obligations(sheet: Worksheet, register: Register, shard: Shard) -> None:
    _header(
        sheet,
        (
            "child_id",
            "public_coordinate",
            "status",
            "derivation",
            "derivation_rationale",
            "obligation_statement",
            "quotes",
            "evidence_types",
            "document_categories",
            "calculation_rules",
            "specialists",
            "independence",
            "acceptance_case",
            "destination",
            "ambiguity_ids",
            "cross_row_instruction_ids",
            "retired_reason",
        ),
    )
    rows = [
        (parent.public_coordinate, child)
        for parent in ordered_parents(register, (shard,))
        for _file, child in register.children
        if child.parent_id == parent.parent_id
    ]
    rows.sort(key=lambda item: (item[0], item[1].child_id))
    for coordinate, child in rows:
        _row(sheet, _obligation_values(register, coordinate, child))


def _ambiguities(sheet: Worksheet, register: Register, shard: Shard) -> None:
    _header(
        sheet,
        (
            "ambiguity_id",
            "topic_code",
            "decision_owner",
            "status",
            "blocks_definitive_answer",
            "affected_coordinates",
            "description",
        ),
    )
    parent_ids = {parent.parent_id for parent in ordered_parents(register, (shard,))}
    rows = []
    seen: set[str] = set()
    for file_shard, item in register.ambiguities:
        named = any(parent_id in parent_ids for parent_id in item.affected_parent_ids)
        if file_shard is not shard and not named:
            continue
        if item.ambiguity_id in seen:
            continue
        seen.add(item.ambiguity_id)
        rows.append(item)
    rows.sort(key=lambda item: item.ambiguity_id)
    for item in rows:
        coordinates = ",".join(
            source_coordinate(register, parent_id) for parent_id in item.affected_parent_ids
        )
        _row(
            sheet,
            [
                item.ambiguity_id,
                item.topic_code.value,
                item.decision_owner.value,
                item.status.value,
                "true" if item.blocks_definitive_answer else "false",
                coordinates,
                item.description,
            ],
        )


def _cross_rows(sheet: Worksheet, register: Register, shard: Shard) -> None:
    _header(
        sheet,
        (
            "instruction_id",
            "kind",
            "source_coordinate",
            "source_field",
            "target_coordinates",
            "note",
            "quote_text",
        ),
    )
    rows = [item for file_shard, item in register.cross_rows if file_shard is shard]
    rows.sort(key=lambda item: item.instruction_id)
    for item in rows:
        targets = ",".join(
            source_coordinate(register, parent_id) for parent_id in item.target_parent_ids
        )
        _row(
            sheet,
            [
                item.instruction_id,
                item.kind.value,
                source_coordinate(register, item.source_parent_id),
                item.source_field,
                targets,
                item.note,
                item.source_quote.text,
            ],
        )


def _dispositions(sheet: Worksheet, register: Register, shard: Shard) -> None:
    _header(
        sheet,
        (
            "kind",
            "public_coordinate",
            "field",
            "start",
            "end",
            "disposition",
            "disposed_text",
            "note",
            "child_ids",
        ),
    )
    rows = [item for file_shard, item in register.dispositions if file_shard is shard]
    rows.sort(
        key=lambda item: (
            source_coordinate(register, item.parent_id),
            item.kind.value,
            item.field,
            item.start,
            item.end,
        )
    )
    for item in rows:
        _row(
            sheet,
            [
                item.kind.value,
                source_coordinate(register, item.parent_id),
                item.field,
                item.start,
                item.end,
                item.disposition,
                disposed_text(register, item),
                item.note or "",
                ",".join(item.child_ids),
            ],
        )


def _obligation_values(register: Register, coordinate: str, child: ChildObligation) -> list[object]:
    quotes = _quote_lines(register, child)
    rules = "\n".join(
        f"kind={rule.kind.value} code={rule.rule_code or ''} note={rule.note or ''}"
        for rule in child.calculation_rules
    )
    specialist = child.responsible_specialist
    names = (
        "" if specialist is None else ",".join(item.value for item in specialist.registered_types)
    )
    destination = child.output_destination
    destination_text = ""
    if destination is not None:
        destination_text = (
            f"{destination.template}:{destination.report_section}:"
            f"{destination.report_subsection or ''}"
        )
    return [
        child.child_id,
        coordinate,
        child.status.value,
        child.derivation.value,
        child.derivation_rationale or "",
        child.obligation_statement,
        quotes,
        ",".join(item.evidence_type.value for item in child.acceptable_evidence),
        ",".join(
            category for item in child.acceptable_evidence for category in item.document_categories
        ),
        rules,
        names,
        ",".join(item.independence.value for item in child.acceptable_evidence),
        child.acceptance_case,
        destination_text,
        ",".join(child.ambiguity_ids),
        ",".join(child.cross_row_instruction_ids),
        child.retired_reason or "",
    ]


def _header(sheet: Worksheet, names: tuple[str, ...]) -> None:
    _row(sheet, list(names))


def _quote_lines(register: Register, child: ChildObligation) -> str:
    lines: list[str] = []
    for quote in child.source_quotes:
        parent = source_coordinate(register, quote.parent_id or child.parent_id)
        lines.append(f"field={quote.field} parent={parent} text={quote.text}")
    return "\n".join(lines)


def _row(sheet: Worksheet, values: list[object]) -> None:
    cells: list[object] = []
    for value in values:
        if isinstance(value, int) and not isinstance(value, bool):
            cells.append(value)
        else:
            cells.append(_cell_text("" if value is None else str(value)))
    sheet.append(cells)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"
            cell.number_format = "@"
        cell.alignment = _WRAP


def _cell_text(value: str) -> str:
    """Replace XML-illegal controls and cap a cell at the worksheet limit."""
    escaped = _escape_controls(value)
    if len(escaped) <= _CELL_MAX_CHARS:
        return escaped
    keep = _CELL_MAX_CHARS - len(_CELL_CUT_SUFFIX)
    return escaped[:keep] + _CELL_CUT_SUFFIX


def _escape_controls(value: str) -> str:
    pieces: list[str] = []
    for character in value:
        codepoint = ord(character)
        if codepoint < _C0_CONTROL_LIMIT and character not in "\t\n\r":
            pieces.append(f"\\u{codepoint:04x}")
        else:
            pieces.append(character)
    return "".join(pieces)


def _finish(sheet: Worksheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index in range(1, sheet.max_column + 1):
        width = _ID_WIDTH if index <= 2 else _TEXT_WIDTH
        sheet.column_dimensions[get_column_letter(index)].width = width
