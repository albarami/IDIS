"""Fictional workbooks and authored registers for GL-02 tests."""

from __future__ import annotations

from collections import defaultdict
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment

from idis.methodology.obligations.extract import extract_register
from idis.methodology.obligations.jsonl import write_jsonl
from idis.methodology.obligations.models_authored import (
    Ambiguity,
    AnalystReview,
    Applicability,
    BaselineExpectation,
    Capability,
    ChildObligation,
    ConsumerDetail,
    Coordinates,
    CoordinateSlot,
    EvidenceRequirement,
    OutputDestination,
    ParentProfile,
    ProducerDetail,
    RequiredFact,
    ResponsibleSpecialist,
    RoleAlias,
    SourceQuote,
    ViewRecord,
)
from idis.methodology.obligations.register_io import load_register, parent_shard
from idis.methodology.obligations.runs import field_text
from idis.methodology.obligations.vocab import (
    NAMED_AMBIGUITY_TOPICS,
    AmbiguityDecisionOwner,
    AmbiguityStatus,
    ApplicabilityDecision,
    BaselineKind,
    CapabilityStatus,
    ChildStatus,
    ConsumerKind,
    CoordinateRequirement,
    Derivation,
    EvidenceType,
    FundFeedbackStatus,
    Independence,
    OutputDestinationStatus,
    ParentKind,
    PlanMinimumCode,
    ProducerKind,
    RoleAliasOwner,
    RoleAliasStatus,
    Shard,
    SpecialistAgentType,
    ViewClosure,
    ViewStatus,
    plan_code_shard,
)
from idis.models.document_classification import CddDocumentCategory
from tests.slice100_task5._workbook_factory import production_fdd_workbook_bytes

_TERM_PAIRS = {
    "P&L": ((4, 5), (6, 7), (8, 9), (10, 11), (12, 13), (14, 15)),
    "Cash Flow": ((4, 5), (6, 7), (8, 9), (10, 11), (12, 13)),
    "Liabilities": ((4, 5), (6, 7), (8, 9), (10, 11), (12, 13), (14, 15)),
    "Assets": ((4, 5), (6, 7), (8, 9), (10, 11)),
}
_NATURE_PAIRS = {
    "P&L": ((16, 17), (18, 19), (20, 21), (22, 23), (24, 25)),
    "Liabilities": ((19, 20), (21, 22)),
    "Assets": ((12, 13), (14, 15), (16, 17)),
}
_LIABILITY_QUESTION_MERGES = ((4, 6), (17, 18))
_CDD_ALLOWED = frozenset(
    item.value for item in CddDocumentCategory if item.value not in {"unsupported", "unknown"}
)


def structural_workbook_bytes() -> bytes:
    """Production-shaped workbook with the published merge counts and cleared merge children."""
    workbook = load_workbook(BytesIO(production_fdd_workbook_bytes()))
    liabilities = workbook["Liabilities"]
    for first_row, last_row in _LIABILITY_QUESTION_MERGES:
        liabilities.unmerge_cells(
            start_row=first_row, start_column=5, end_row=last_row, end_column=5
        )
        for row_index in range(first_row + 1, last_row + 1):
            liabilities.cell(row=row_index, column=5).value = None
        liabilities.merge_cells(start_row=first_row, start_column=5, end_row=last_row, end_column=5)
    for sheet_name, pairs in _TERM_PAIRS.items():
        _merge_pairs(workbook[sheet_name], pairs, 2)
    for sheet_name, pairs in _NATURE_PAIRS.items():
        _merge_pairs(workbook[sheet_name], pairs, 3)
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def feature_workbook_bytes() -> bytes:
    """Structural workbook plus a comment, a hidden row, a non-string and an outside cell."""
    workbook = load_workbook(BytesIO(structural_workbook_bytes()))
    sheet = workbook["P&L"]
    outside = sheet.cell(row=1, column=1)
    outside.value = "Fictional outside cell"
    outside.comment = Comment("Fictional comment", "Fictional author")
    sheet.row_dimensions[1].hidden = True
    sheet.cell(row=40, column=1).value = 7
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def whitespace_workbook_bytes() -> bytes:
    """Structural workbook with edge whitespace and a line break in four roles."""
    workbook = load_workbook(BytesIO(structural_workbook_bytes()))
    sheet = workbook["P&L"]
    sheet.cell(row=26, column=2).value = " fictional-P&L-term-23 "
    sheet.cell(row=26, column=3).value = " fictional-P&L-nature-23 "
    sheet.cell(row=26, column=4).value = " fictional-P&L-line-23\n "
    sheet.cell(row=26, column=5).value = " Review fictional P&L evidence row 23. "
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def build_authored_register(directory: Path, *, policy: Path | None = None) -> tuple[Path, Path]:
    """Extract a structural workbook and author a register that satisfies every local rule."""
    directory.mkdir(parents=True, exist_ok=True)
    root = directory / "register"
    workbook_path = directory / "fictional.xlsx"
    workbook_path.write_bytes(structural_workbook_bytes())
    extract_register(workbook_path, root, policy_path=policy)
    _author(root)
    return root, workbook_path


def _merge_pairs(sheet, pairs: tuple[tuple[int, int], ...], column: int) -> None:
    for first_row, last_row in pairs:
        for row_index in range(first_row + 1, last_row + 1):
            sheet.cell(row=row_index, column=column).value = None
        sheet.merge_cells(
            start_row=first_row, start_column=column, end_row=last_row, end_column=column
        )


def _author(root: Path) -> None:
    register = load_register(root)
    codes_by_shard: dict[Shard, list[str]] = defaultdict(list)
    for code in PlanMinimumCode:
        codes_by_shard[plan_code_shard(code.value)].append(code.value)
    seen: set[Shard] = set()
    profiles: dict[Shard, list[ParentProfile]] = defaultdict(list)
    children: dict[Shard, list[ChildObligation]] = defaultdict(list)
    roles: set[str] = set()
    for parent in register.parents:
        shard = parent_shard(parent)
        profiles[shard].append(_profile(parent))
        plan_codes: tuple[str, ...] = ()
        if parent.parent_kind is ParentKind.FINANCIAL and shard not in seen:
            seen.add(shard)
            plan_codes = tuple(codes_by_shard[shard])
        children[shard].append(_child(parent, plan_codes))
        if parent.commercial_context is not None:
            roles.update(parent.commercial_context.template_role)
    for shard, records in profiles.items():
        write_jsonl(
            root / "authored" / "profiles" / f"{shard.value}.jsonl",
            records,
            key=lambda item: item.parent_id,
        )
    for shard, records in children.items():
        write_jsonl(
            root / "authored" / "children" / f"{shard.value}.jsonl",
            records,
            key=lambda item: item.child_id,
        )
    financial = next(
        parent for parent in register.parents if parent.parent_kind is ParentKind.FINANCIAL
    )
    ambiguities = [
        Ambiguity(
            ambiguity_id=f"AMB-{topic.value}",
            topic_code=topic,
            affected_parent_ids=(financial.parent_id,),
            description=f"Fictional ambiguity {topic.value}.",
            decision_owner=AmbiguityDecisionOwner.FUND,
            blocks_definitive_answer=False,
            status=AmbiguityStatus.OPEN,
        )
        for topic in NAMED_AMBIGUITY_TOPICS
    ]
    write_jsonl(
        root / "authored" / "ambiguities" / "XS.jsonl",
        ambiguities,
        key=lambda item: item.ambiguity_id,
    )
    aliases = [
        RoleAlias(
            template_role=role,
            proposed_registered_types=(SpecialistAgentType.FINANCIAL_AGENT,),
            rationale="Fictional role alias.",
            decision_owner=RoleAliasOwner.GL_12,
            status=RoleAliasStatus.PROPOSED,
        )
        for role in sorted(roles)
    ]
    write_jsonl(
        root / "authored" / "role_aliases" / "CDD.jsonl",
        aliases,
        key=lambda item: item.template_role,
    )
    write_jsonl(
        root / "authored" / "baseline" / "XS.jsonl",
        [
            BaselineExpectation(
                expectation_id="BASE-LIAB-004",
                kind=BaselineKind.MERGED_INSTRUCTION,
                source_coordinate="FDD-LIAB-004",
                target_coordinates=(),
                all_rows_of_sheet=None,
                note="Fictional merged instruction baseline.",
            )
        ],
        key=lambda item: item.expectation_id,
    )
    write_jsonl(
        root / "authored" / "views" / "XS.jsonl",
        [
            ViewRecord(
                view_id="fictional-pl",
                parent_ids=("FDD-PL-004",),
                closure=ViewClosure.NONE,
                status=ViewStatus.PROPOSED,
                agreed_by=None,
            )
        ],
        key=lambda item: item.view_id,
    )


def _profile(parent) -> ParentProfile:
    template_role = None
    if parent.commercial_context is not None and parent.commercial_context.template_role:
        template_role = parent.commercial_context.template_role[0]
    return ParentProfile(
        parent_id=parent.parent_id,
        default_coordinates=_coordinates(),
        default_producer=ProducerDetail(
            kind=ProducerKind.METHODOLOGY_EXTRACTION, detail="Fictional producer."
        ),
        default_consumer=ConsumerDetail(
            kind=ConsumerKind.REPORT_SECTION, detail="Fictional consumer."
        ),
        default_specialists=ResponsibleSpecialist(
            registered_types=(SpecialistAgentType.FINANCIAL_AGENT,),
            template_role=template_role,
        ),
        default_output_destination=_destination(parent),
        default_analyst_review=AnalystReview(required=True, reason="Fictional analyst review."),
        default_capability=Capability(status=CapabilityStatus.UNASSESSED),
        default_fund_feedback_status=FundFeedbackStatus.PROPOSED,
        parent_summary_note="Fictional parent summary.",
    )


def _destination(parent) -> OutputDestination:
    if parent.commercial_context is not None:
        mapping = parent.commercial_context.report_mapping
        return OutputDestination(
            template="commercial_dd",
            report_section=mapping.report_section,
            report_subsection=mapping.report_subsection,
            status=OutputDestinationStatus.PROVISIONAL,
        )
    return OutputDestination(
        template="financial_dd",
        report_section=parent.sheet or "P&L",
        report_subsection=parent.line_item_raw or "provisional",
        status=OutputDestinationStatus.PROVISIONAL,
    )


def _child(parent, plan_codes: tuple[str, ...]) -> ChildObligation:
    commercial = parent.parent_kind is ParentKind.COMMERCIAL
    quotes = []
    for field_name in ("question", "line_item", "term", "nature"):
        text = field_text(parent, field_name)
        if text:
            quotes.append(SourceQuote(field=field_name, occurrence=0, parent_id=None, text=text))
    categories: tuple[str, ...] = ("pl_support",)
    if commercial and parent.commercial_context is not None:
        chosen = next(
            (
                item
                for item in parent.commercial_context.governed_categories
                if item in _CDD_ALLOWED
            ),
            None,
        )
        categories = () if chosen is None else (chosen,)
    return ChildObligation(
        acceptable_evidence=(
            EvidenceRequirement(
                description="Fictional ledger extract.",
                document_categories=categories,
                evidence_type=EvidenceType.LEDGER_EXTRACT,
                independence=Independence.INTERNAL,
                min_count=1,
            ),
        ),
        acceptance_case="Fictional acceptance case.",
        applicability=Applicability(condition=None, decision=ApplicabilityDecision.ALWAYS),
        child_id=f"GL02-{parent.public_coordinate}-C001",
        completeness_conditions=("Fictional completeness condition.",),
        derivation=Derivation.TEMPLATE_FIELD if commercial else Derivation.EXPLICIT,
        derivation_rationale="Fictional template rationale." if commercial else None,
        joint_answer_reason="Fictional joint answer reason.",
        judgment_sought="Fictional judgment sought." if commercial else None,
        obligation_statement="Fictional obligation statement.",
        parent_id=parent.parent_id,
        plan_minimum_codes=plan_codes,
        required_facts=(
            RequiredFact(
                description="Fictional fact.",
                fact_id="F01",
                needs=("entity", "period", "currency", "unit", "adjustment"),
            ),
        ),
        shared_instruction_ids=parent.merged_instruction_ids,
        source_quotes=tuple(quotes),
        status=ChildStatus.ACTIVE,
    )


def _coordinates() -> Coordinates:
    slot = CoordinateSlot(requirement=CoordinateRequirement.REQUIRED, note=None)
    return Coordinates(
        entity=slot,
        consolidation_scope=slot,
        period=slot,
        actual_forecast_status=slot,
        currency=slot,
        unit=slot,
        adjustment=slot,
    )
