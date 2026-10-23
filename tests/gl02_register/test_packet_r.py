"""Views, review pack, audit input, reference resolution and the two real-data defects."""

from __future__ import annotations

import importlib.util
import json
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from idis.methodology.obligations.audit_input import command_audit_input
from idis.methodology.obligations.commands import command_views
from idis.methodology.obligations.errors import Gl02Error
from idis.methodology.obligations.jsonl import read_jsonl, write_json, write_jsonl
from idis.methodology.obligations.models_authored import (
    CalculationRule,
    ChildObligation,
    ViewRecord,
)
from idis.methodology.obligations.register_io import load_register, parent_shard
from idis.methodology.obligations.review_pack import command_review_pack
from idis.methodology.obligations.validate import invariant_status, validate_register
from idis.methodology.obligations.vocab import (
    SKIPPED_LEDGER_ABSENT,
    CalculationKind,
    ChildStatus,
    Code,
    FindingSeverity,
    InvariantStatus,
)
from tests.gl02_register.mutate import clone_register, finding_codes, update_child
from tests.gl02_register.support import build_authored_register, structural_workbook_bytes

_MARKERS = ("=1+1", "+1", "-1", "@cmd")


def test_validate_json_includes_skipped_invariants(
    authored_register: tuple[Path, Path], capsys: pytest.CaptureFixture[str]
) -> None:
    root, _workbook = authored_register
    cli = _cli()
    assert cli.main(["validate", "--private-root", str(root), "--format", "json"]) in {0, 1}
    payload = json.loads(capsys.readouterr().out)
    pairs = payload["skipped_invariants"]
    assert pairs == sorted(pairs, key=lambda item: (item["invariant"], item["code"]))
    assert {"code": SKIPPED_LEDGER_ABSENT, "invariant": "I-16"} in pairs
    assert "skipped_checks" in payload


def test_validate_view_limits_findings(
    authored_register: tuple[Path, Path], tmp_path: Path
) -> None:
    root = clone_register(authored_register[0], tmp_path / "view")
    _drop_child(root, "FDD-PL-004")
    _drop_child(root, "FDD-CF-004")
    register = load_register(root)
    inside = register.parents_by_coordinate["FDD-PL-004"].parent_id
    outside = register.parents_by_coordinate["FDD-CF-004"].parent_id
    report = validate_register(root, view="fictional-pl")
    assert _has_parent(report, Code.GL02_PARENT_WITHOUT_CHILD.value, inside)
    assert not _has_parent(report, Code.GL02_PARENT_WITHOUT_CHILD.value, outside)


def test_validate_unknown_view(authored_register: tuple[Path, Path]) -> None:
    root, _workbook = authored_register
    report = validate_register(root, view="fictional-missing")
    assert any(
        item.code == Code.GL02_VIEW_INVALID.value and "fictional-missing" in item.ids
        for item in report.findings
    )


def test_views_command_hides_the_name(authored_register: tuple[Path, Path], tmp_path: Path) -> None:
    root = clone_register(authored_register[0], tmp_path / "views")
    path = root / "authored" / "views" / "XS.jsonl"
    records = [
        item.model_copy(update={"agreed_by": "Fictional Reviewer"})
        for item in read_jsonl(path, ViewRecord)
    ]
    write_jsonl(path, records, key=lambda item: item.view_id)
    text = command_views(root)
    payload = json.loads(command_views(root, output_format="json"))
    assert "Fictional Reviewer" not in text
    assert "Fictional Reviewer" not in json.dumps(payload)
    row = payload["views"][0]
    assert row["view_id"] == "fictional-pl"
    assert row["agreed_by_present"] is True
    assert row["closure"] == "none"
    assert row["primary_parent_count"] == 1
    assert row["closure_coordinates"] == ["FDD-PL-004"]


def test_review_pack_is_private_ordered_and_shard_filtered(
    authored_register: tuple[Path, Path], tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    root = clone_register(authored_register[0], tmp_path / "pack")
    outside = tmp_path / "outside"
    summary = command_review_pack(root, out_dir=root / "review_pack")
    assert "Review fictional" not in summary
    assert "Fictional obligation statement." not in summary
    markdown = (root / "review_pack" / "PL" / "obligations.md").read_text(encoding="utf-8")
    assert markdown.index("## FDD-PL-004") < markdown.index("## FDD-PL-005")
    assert "Fictional obligation statement." in markdown
    assert (root / "review_pack" / "PL.xlsx").is_file()
    assert (root / "review_pack" / "README.md").is_file()
    filtered = command_review_pack(root, out_dir=root / "only-pl", shards=["PL"])
    assert "CF/obligations.md" not in filtered
    assert "PL/obligations.md" in filtered
    assert not (root / "only-pl" / "CF").exists()
    with pytest.raises(Gl02Error):
        command_review_pack(root, out_dir=outside)
    assert not outside.exists()
    cli = _cli()
    assert cli.main(["review-pack", "--private-root", str(root), "--out-dir", str(outside)]) == 2
    captured = capsys.readouterr()
    assert "Review fictional" not in captured.out
    assert "Fictional obligation statement." not in captured.out


def test_review_pack_stores_formula_text_as_strings(
    authored_register: tuple[Path, Path], tmp_path: Path
) -> None:
    root = clone_register(authored_register[0], tmp_path / "inject")
    update_child(root, "FDD-PL-004", _inject)
    command_review_pack(root, out_dir=root / "review_pack", shards=["PL"])
    path = root / "review_pack" / "PL.xlsx"
    found: list[str] = []
    workbook = load_workbook(path)
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.data_type != "f"
                    if isinstance(cell.value, str) and cell.value[:1] in "=+-@":
                        found.append(cell.value)
                        assert cell.data_type == "s"
    finally:
        workbook.close()
    for marker in _MARKERS:
        assert marker in found
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith("xl/worksheets/"):
                xml = archive.read(name)
                assert b"<f>" not in xml
                assert b"<f " not in xml


def test_audit_input_blind_omits_authored_obligations(
    authored_register: tuple[Path, Path], tmp_path: Path
) -> None:
    root = clone_register(authored_register[0], tmp_path / "audit")
    blind_out = command_audit_input(root, stage="blind", shard="PL")
    blind = (root / "omission_audit_input" / "PL-blind.md").read_text(encoding="utf-8")
    compare_out = command_audit_input(root, stage="compare", shard="PL")
    compare = (root / "omission_audit_input" / "PL-compare.md").read_text(encoding="utf-8")
    assert "Fictional obligation statement." not in blind_out
    assert "Fictional obligation statement." not in compare_out
    register = load_register(root)
    for _shard, child in register.children:
        if parent_shard(register.parent(child.parent_id) or register.parents[0]) is not None:
            pass
        if child.child_id.startswith("GL02-FDD-PL-"):
            assert child.child_id not in blind
            assert child.obligation_statement not in blind
            assert child.child_id in compare
            assert child.obligation_statement in compare
    assert blind.index("## FDD-PL-004") < blind.index("## FDD-PL-005")
    assert "plan_minimum_codes=" in blind
    again = root / "omission_audit_input" / "again.md"
    command_audit_input(root, stage="blind", shard="PL", out_file=again)
    assert again.read_bytes() == blind.encode("utf-8")
    with pytest.raises(Gl02Error):
        command_audit_input(root, stage="blind", shard="PL", out_file=tmp_path / "outside.md")


def test_derived_rule_scope_passes(authored_register: tuple[Path, Path], tmp_path: Path) -> None:
    root = clone_register(authored_register[0], tmp_path / "scope-pass")
    _set_rules(root, "FDD-PL-004", ("RETAINED_EARNINGS",))
    commercial = _commercial(root)
    _set_rules(root, commercial, ("CAGR", "RETAINED_EARNINGS"))
    assert Code.GL02_DERIVED_RULE_SCOPE.value not in finding_codes(root)


def test_derived_rule_scope_mutation(authored_register: tuple[Path, Path], tmp_path: Path) -> None:
    root = clone_register(authored_register[0], tmp_path / "scope-fail")
    before = finding_codes(root)
    _set_rules(root, "FDD-PL-004", ("CAGR",))
    assert finding_codes(root) - before == {Code.GL02_DERIVED_RULE_SCOPE.value}
    report = validate_register(root)
    assert any(
        item.code == Code.GL02_DERIVED_RULE_SCOPE.value and "CAGR" in item.ids
        for item in report.findings
    )


def test_policy_given_not_extracted(tmp_path: Path) -> None:
    root, _workbook = build_authored_register(tmp_path / "given")
    policy = tmp_path / "given" / "fictional-policy.json"
    policy.write_text("{}\n", encoding="utf-8", newline="\n")
    report = validate_register(root, policy_path=policy)
    assert Code.GL02_POLICY_GIVEN_NOT_EXTRACTED.value in {item.code for item in report.findings}
    assert Code.GL02_POLICY_UNCARRIED.value not in {item.code for item in report.findings}
    assert invariant_status("I-10", report) == InvariantStatus.FAIL.value
    assert "I-10" not in {item_id for item_id, _code in report.skipped_invariants}


def test_policy_coverage_is_shard_limited(tmp_path: Path) -> None:
    root, policy = _two_shard_policy(tmp_path)
    register = load_register(root)
    carried = register.parents_by_coordinate["FDD-PL-004"].parent_id
    missing = register.parents_by_coordinate["FDD-CF-004"].parent_id
    full = validate_register(root, policy_path=policy)
    assert _has_parent(full, Code.GL02_POLICY_UNCARRIED.value, missing)
    assert not _has_parent(full, Code.PENDING_CROSS_SHARD.value, missing)
    limited = validate_register(root, policy_path=policy, shards=["PL"])
    assert _has_parent(limited, Code.PENDING_CROSS_SHARD.value, missing)
    assert not _has_parent(limited, Code.GL02_POLICY_UNCARRIED.value, missing)
    assert not _has_parent(limited, Code.GL02_POLICY_UNCARRIED.value, carried)
    owned = validate_register(root, policy_path=policy, shards=["CF"])
    assert _has_parent(owned, Code.GL02_POLICY_UNCARRIED.value, missing)
    pending = [
        item
        for item in limited.findings
        if item.code == Code.PENDING_CROSS_SHARD.value and missing in item.ids
    ]
    errors = [
        item
        for item in full.findings
        if item.code == Code.GL02_POLICY_UNCARRIED.value and missing in item.ids
    ]
    assert pending and errors
    assert pending[0].ids == errors[0].ids
    assert pending[0].severity is FindingSeverity.INFO


def test_i19_references_pass(authored_register: tuple[Path, Path]) -> None:
    root, _workbook = authored_register
    assert Code.GL02_REFERENCE_UNRESOLVED.value not in finding_codes(root)
    assert Code.GL02_REFERENCE_RETIRED.value not in finding_codes(root)
    assert invariant_status("I-19", validate_register(root)) == InvariantStatus.PASS.value


def test_i19_reference_unresolved(authored_register: tuple[Path, Path], tmp_path: Path) -> None:
    root = clone_register(authored_register[0], tmp_path / "missing-ref")
    _cite(root, "FDD-PL-004", "Fictional obligation GL02-FDD-PL-999-C001.")
    report = validate_register(root)
    assert any(
        item.code == Code.GL02_REFERENCE_UNRESOLVED.value and "GL02-FDD-PL-999-C001" in item.ids
        for item in report.findings
    )


def test_i19_reference_retired(authored_register: tuple[Path, Path], tmp_path: Path) -> None:
    root = clone_register(authored_register[0], tmp_path / "retired-ref")
    _add_retired(root, "FDD-PL-004")
    _cite(root, "FDD-PL-004", "Fictional obligation GL02-FDD-PL-004-C002.")
    report = validate_register(root)
    assert any(
        item.code == Code.GL02_REFERENCE_RETIRED.value
        and item.severity is FindingSeverity.WARNING
        and "GL02-FDD-PL-004-C002" in item.ids
        for item in report.findings
    )
    assert not any(
        item.code == Code.GL02_REFERENCE_UNRESOLVED.value and "GL02-FDD-PL-004-C002" in item.ids
        for item in report.findings
    )


def test_i19_reference_unknown_coordinate(
    authored_register: tuple[Path, Path], tmp_path: Path
) -> None:
    root = clone_register(authored_register[0], tmp_path / "coord")
    _cite(root, "FDD-PL-004", "Fictional obligation FDD-PL-999.")
    _cite(root, "FDD-PL-005", "Fictional obligation GL02-FDD-PL-999-C001Z.")
    report = validate_register(root)
    assert any(
        item.code == Code.GL02_REFERENCE_UNRESOLVED.value and "FDD-PL-999" in item.ids
        for item in report.findings
    )
    assert not any("GL02-FDD-PL-999-C001Z" in item.ids for item in report.findings)


def test_i19_reference_outside_shard_is_pending(
    authored_register: tuple[Path, Path], tmp_path: Path
) -> None:
    root = clone_register(authored_register[0], tmp_path / "shard-ref")
    _cite(root, "FDD-PL-004", "Fictional obligation FDD-CF-004.")
    full = validate_register(root)
    assert not any(
        item.code == Code.GL02_REFERENCE_UNRESOLVED.value and "FDD-CF-004" in item.ids
        for item in full.findings
    )
    limited = validate_register(root, shards=["PL"])
    pending = [
        item
        for item in limited.findings
        if item.code == Code.PENDING_CROSS_SHARD.value and "FDD-CF-004" in item.ids
    ]
    assert pending
    assert all(item.severity is FindingSeverity.INFO for item in pending)
    assert not any(
        item.code == Code.GL02_REFERENCE_UNRESOLVED.value and "FDD-CF-004" in item.ids
        for item in limited.findings
    )


def _cli():
    spec = importlib.util.spec_from_file_location(
        "gl02_register_cli_r", Path("scripts/gl02_register.py")
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _has_parent(report: object, code: str, parent_id: str) -> bool:
    findings = report.findings
    return any(item.code == code and parent_id in item.ids for item in findings)


def _drop_child(root: Path, coordinate: str) -> None:
    register = load_register(root)
    parent = register.parents_by_coordinate[coordinate]
    shard = parent_shard(parent)
    path = root / "authored" / "children" / f"{shard.value}.jsonl"
    children = [
        child for child in read_jsonl(path, ChildObligation) if child.parent_id != parent.parent_id
    ]
    write_jsonl(path, children, key=lambda item: item.child_id)


def _cite(root: Path, coordinate: str, statement: str) -> None:
    update_child(
        root,
        coordinate,
        lambda child: child.model_copy(update={"obligation_statement": statement}),
    )


def _inject(child: ChildObligation) -> ChildObligation:
    return child.model_copy(
        update={
            "obligation_statement": "=1+1",
            "acceptance_case": "+1",
            "derivation_rationale": "-1",
            "retired_reason": "@cmd",
        }
    )


def _set_rules(root: Path, coordinate: str, codes: tuple[str, ...]) -> None:
    rules = tuple(
        CalculationRule(kind=CalculationKind.DERIVED_GL08, rule_code=code, note=None)
        for code in codes
    )
    update_child(
        root,
        coordinate,
        lambda child: child.model_copy(update={"calculation_rules": rules}),
    )


def _commercial(root: Path) -> str:
    register = load_register(root)
    return next(
        parent.public_coordinate
        for parent in register.parents
        if parent.public_coordinate.startswith("CDD-")
    )


def _add_retired(root: Path, coordinate: str) -> None:
    register = load_register(root)
    parent = register.parents_by_coordinate[coordinate]
    shard = parent_shard(parent)
    path = root / "authored" / "children" / f"{shard.value}.jsonl"
    children = list(read_jsonl(path, ChildObligation))
    source = next(child for child in children if child.parent_id == parent.parent_id)
    children.append(
        source.model_copy(
            update={
                "child_id": f"GL02-{coordinate}-C002",
                "status": ChildStatus.RETIRED,
                "retired_reason": "Fictional retirement.",
                "obligation_statement": "Fictional retired obligation.",
            }
        )
    )
    write_jsonl(path, children, key=lambda item: item.child_id)


def _two_shard_policy(directory: Path) -> tuple[Path, Path]:
    workbook_bytes = structural_workbook_bytes()
    preview, _workbook = build_authored_register(
        directory / "preview", workbook_bytes=workbook_bytes
    )
    register = load_register(preview)
    pl_parent = register.parents_by_coordinate["FDD-PL-004"].parent_id
    cf_parent = register.parents_by_coordinate["FDD-CF-004"].parent_id
    policy = directory / "fictional-policy.json"
    write_json(
        policy,
        {
            "workbook_row_registry": {
                "rows": [
                    _policy_row("P&L", 4, pl_parent),
                    _policy_row("Cash Flow", 4, cf_parent),
                ]
            },
            "formula_registry": {
                "rows": [
                    _formula(1, "P&L", 4, "P&L", 4),
                    _formula(2, "Cash Flow", 4, "Cash Flow", 4),
                ]
            },
        },
    )
    root, _workbook = build_authored_register(
        directory / "register", policy=policy, workbook_bytes=workbook_bytes
    )
    update_child(
        root,
        "FDD-PL-004",
        lambda child: child.model_copy(
            update={
                "calculation_rules": (
                    CalculationRule(kind=CalculationKind.POLICY_V2_ANCHOR, formula_ordinal=1),
                    CalculationRule(
                        kind=CalculationKind.POLICY_V2_BINDING,
                        formula_ordinal=1,
                        binding_ordinal=1,
                    ),
                )
            }
        ),
    )
    return root, policy


def _policy_row(sheet: str, row: int, parent_id: str) -> dict[str, object]:
    return {
        "methodology_question_id": parent_id,
        "required_calculations": [],
        "source_row": row,
        "source_sheet": sheet,
    }


def _formula(
    ordinal: int, anchor_sheet: str, anchor_row: int, source_sheet: str, source_row: int
) -> dict[str, object]:
    return {
        "anchor_row": anchor_row,
        "anchor_sheet": anchor_sheet,
        "formula_ordinal": ordinal,
        "formula_type": "fictional",
        "ordered_input_specs": [
            {
                "binding_kind": "SOURCE",
                "binding_ordinal": 1,
                "cardinality": "1",
                "period_role": None,
                "predicate": None,
                "role": "fictional",
                "sign_transform": None,
                "source_row": source_row,
                "source_sheet": source_sheet,
            }
        ],
        "required": True,
        "scope": "fictional",
    }
