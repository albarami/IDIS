"""Corrections from the packet R review. Fixtures are fictional."""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest
from openpyxl import load_workbook

from idis.methodology.obligations.audit_input import command_audit_input
from idis.methodology.obligations.errors import Gl02Error
from idis.methodology.obligations.jsonl import read_jsonl, write_jsonl
from idis.methodology.obligations.ledger import command_seal
from idis.methodology.obligations.models_authored import (
    ChildObligation,
    CrossRowInstruction,
    LedgerRecord,
    SourceQuote,
)
from idis.methodology.obligations.register_io import load_register
from idis.methodology.obligations.review_pack import command_review_pack
from idis.methodology.obligations.review_pack_xlsx import _CELL_CUT_SUFFIX, _CELL_MAX_CHARS
from idis.methodology.obligations.validate import invariant_status, validate_register
from idis.methodology.obligations.vocab import (
    Code,
    CrossRowKind,
    FindingSeverity,
    InvariantStatus,
)
from tests.gl02_register.mutate import clone_register, finding_codes, update_child, update_profile
from tests.gl02_register.test_invariants_later import _policy_register
from tests.gl02_register.test_ledger import (
    _NOW,
    _NOW_2,
    _add_child,
    _drop_child,
    _files,
    _pin_fictional,
    _sealed,
)
from tests.gl02_register.test_packet_r import _two_shard_policy

_MARKER = "FictionalMarkerGL02R"
_CHILD = "GL02-FDD-PL-005-C001"


def test_full_seal_succeeds_and_a_later_edit_appends_one_id(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _policy_register(tmp_path)
    _pin_fictional(monkeypatch, root)
    workbook = root.parent / "fictional.xlsx"
    policy = tmp_path / "fictional-policy.json"
    code, text = command_seal(root, workbook=workbook, policy=policy, now=_NOW)
    assert code == 0, text
    assert text.startswith("seq=1 ")
    after = validate_register(root, workbook_path=workbook, policy_path=policy)
    assert after.register_complete is True
    assert invariant_status("I-16", after) == InvariantStatus.PASS.value
    code, unchanged = command_seal(root, workbook=workbook, policy=policy, now=_NOW_2)
    assert code == 0
    assert unchanged == "UNCHANGED seq=1\n"
    issued = (root / "ledger" / "issued_ids.jsonl").read_bytes()
    _add_child(root, "FDD-PL-004")
    code, edited = command_seal(root, workbook=workbook, policy=policy, now="2026-10-21T01:02:05Z")
    assert code == 0, edited
    assert edited.startswith("seq=2 ")
    assert "new_ids=1 " in edited
    assert (root / "ledger" / "issued_ids.jsonl").read_bytes().startswith(issued)


def test_shard_seal_refuses_an_in_shard_error_and_writes_nothing(
    authored_register: tuple[Path, Path], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = clone_register(authored_register[0], tmp_path / "in-shard")
    _pin_fictional(monkeypatch, root)
    _drop_child(root, "FDD-PL-005")
    before = _files(root)
    code, text = command_seal(root, shards=["PL"], now=_NOW)
    assert code == 1
    assert text.startswith(f"{Code.GL02_MARKER_UNACCOUNTED.value} ")
    assert _files(root) == before
    assert not (root / "ledger").exists()
    assert not (root.parent / f"{root.name}-snapshots").exists()


def test_shard_seal_refuses_an_unowned_error_and_writes_nothing(
    authored_register: tuple[Path, Path], tmp_path: Path
) -> None:
    root = clone_register(authored_register[0], tmp_path / "unowned")
    before = _files(root)
    code, text = command_seal(root, shards=["PL"], now=_NOW)
    assert code == 1
    assert Code.GL02_SOURCE_PIN_MISMATCH.value in text
    assert _files(root) == before
    assert not (root / "ledger").exists()
    assert not (root.parent / f"{root.name}-snapshots").exists()


def test_write_seal_refuses_a_disagreeing_ledger(
    authored_register: tuple[Path, Path], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _sealed(authored_register[0], tmp_path / "disagree", monkeypatch)
    path = root / "ledger" / "issued_ids.jsonl"
    lines = path.read_text(encoding="utf-8").splitlines()
    path.write_text("\n".join(lines[1:]) + "\n", encoding="utf-8", newline="\n")
    before = _files(root)
    monkeypatch.setattr(
        "idis.methodology.obligations.ledger.seal_block_reason",
        lambda *_args, **_kwargs: None,
    )
    code, text = command_seal(root, shards=["PL"], now=_NOW_2)
    assert code == 1
    assert Code.GL02_LEDGER_FILE_INVALID.value in text
    assert _files(root) == before


def test_write_seal_refuses_an_invalid_ledger(
    authored_register: tuple[Path, Path], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _sealed(authored_register[0], tmp_path / "invalid", monkeypatch)
    path = root / "ledger" / "issued_ids.jsonl"
    path.write_bytes(path.read_bytes() + b"not-json\n")
    before = _files(root)
    monkeypatch.setattr(
        "idis.methodology.obligations.ledger.seal_block_reason",
        lambda *_args, **_kwargs: None,
    )
    code, text = command_seal(root, shards=["PL"], now=_NOW_2)
    assert code == 1
    assert Code.GL02_LEDGER_FILE_INVALID.value in text
    assert _files(root) == before


def test_i16_seal_count_mismatch_is_file_invalid(
    authored_register: tuple[Path, Path], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    root = _sealed(authored_register[0], tmp_path / "counts", monkeypatch)
    before = finding_codes(root)
    path = root / "ledger" / "issued_ids.jsonl"
    kept = [item for item in read_jsonl(path, LedgerRecord) if item.child_id != _CHILD]
    write_jsonl(path, kept, key=lambda item: (item.first_seal, item.child_id))
    _drop_child(root, "FDD-PL-005")
    assert Code.GL02_LEDGER_ID_MISSING.value not in finding_codes(root)
    assert finding_codes(root) - before == {
        Code.GL02_LEDGER_FILE_INVALID.value,
        Code.GL02_LEDGER_TREE_CHANGED.value,
        Code.GL02_MARKER_UNACCOUNTED.value,
        Code.GL02_PARENT_WITHOUT_CHILD.value,
        Code.GL02_RUN_UNDISPOSITIONED.value,
    }


def test_review_pack_escapes_controls_and_caps_cells(
    authored_register: tuple[Path, Path],
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    root = clone_register(authored_register[0], tmp_path / "cells")
    update_child(root, "FDD-PL-004", _marked_child)
    summary = command_review_pack(root, out_dir=root / "review_pack", shards=["PL"])
    assert _MARKER not in summary
    sheet = load_workbook(root / "review_pack" / "PL.xlsx")["Obligations"]
    rows = list(sheet.iter_rows(min_row=2))
    row = next(item for item in rows if item[0].value == "GL02-FDD-PL-004-C001")
    statement = row[5].value
    acceptance = row[12].value
    assert isinstance(statement, str)
    assert isinstance(acceptance, str)
    assert "\\u0001" in statement
    assert "\x01" not in statement
    assert "\t" in statement
    assert "\n" in statement
    assert len(acceptance) == _CELL_MAX_CHARS
    assert acceptance.endswith(_CELL_CUT_SUFFIX)
    monkeypatch.setattr(
        "idis.methodology.obligations.review_pack.write_shard_workbook",
        _raising_writer,
    )
    code = _cli().main(
        ["review-pack", "--private-root", str(root), "--out-dir", str(root / "again")]
    )
    captured = capsys.readouterr()
    assert code == 2
    assert "RuntimeError" in captured.err
    assert Code.GL02_IO.value in captured.err
    assert _MARKER not in captured.out
    assert _MARKER not in captured.err
    assert "Traceback" not in captured.err


def test_review_pack_refuses_an_out_dir_inside_the_hashed_tree(
    authored_register: tuple[Path, Path], tmp_path: Path
) -> None:
    root = clone_register(authored_register[0], tmp_path / "confined")
    before = _files(root)
    for name in ("derived", "authored", "ledger"):
        with pytest.raises(Gl02Error) as error:
            command_review_pack(root, out_dir=root / name / "pack")
        assert error.value.code == Code.GL02_USAGE.value
    assert (
        _cli().main(
            ["review-pack", "--private-root", str(root), "--out-dir", str(root / "derived")]
        )
        == 2
    )
    assert _files(root) == before


def test_blind_audit_omits_the_cross_row_note(
    authored_register: tuple[Path, Path], tmp_path: Path
) -> None:
    root = clone_register(authored_register[0], tmp_path / "note")
    register = load_register(root)
    source = register.parents_by_coordinate["FDD-PL-004"]
    target = register.parents_by_coordinate["FDD-PL-005"]
    note = f"Fictional note {_CHILD} {_MARKER}"
    write_jsonl(
        root / "authored" / "cross_row" / "PL.jsonl",
        [
            CrossRowInstruction(
                instruction_id="XS-FICTIONAL-01",
                kind=CrossRowKind.SAME_AS_ABOVE,
                source_parent_id=source.parent_id,
                source_field="question",
                source_quote=SourceQuote(field="question", occurrence=0, text="Fictional quote"),
                target_parent_ids=(target.parent_id,),
                note=note,
            )
        ],
        key=lambda item: item.instruction_id,
    )
    blind_out = command_audit_input(root, stage="blind", shard="PL")
    blind = (root / "omission_audit_input" / "PL-blind.md").read_text(encoding="utf-8")
    compare_out = command_audit_input(root, stage="compare", shard="PL")
    compare = (root / "omission_audit_input" / "PL-compare.md").read_text(encoding="utf-8")
    assert _MARKER not in blind_out
    assert _MARKER not in compare_out
    assert _MARKER not in blind
    assert _CHILD not in blind
    assert "Fictional quote" not in blind
    assert "kind=same_as_above" in blind
    assert "source=FDD-PL-004" in blind
    assert "targets=FDD-PL-005" in blind
    assert _MARKER in compare
    assert note in compare
    assert "Fictional quote" in compare


def test_i19_scans_the_profile_summary_note(
    authored_register: tuple[Path, Path], tmp_path: Path
) -> None:
    root = clone_register(authored_register[0], tmp_path / "summary")
    before = finding_codes(root)
    update_profile(
        root,
        "FDD-PL-004",
        lambda profile: profile.model_copy(
            update={"parent_summary_note": "Fictional summary GL02-FDD-PL-999-C001."}
        ),
    )
    assert finding_codes(root) - before == {Code.GL02_REFERENCE_UNRESOLVED.value}
    report = validate_register(root)
    assert any(
        item.code == Code.GL02_REFERENCE_UNRESOLVED.value and "GL02-FDD-PL-999-C001" in item.ids
        for item in report.findings
    )


def test_limited_run_marks_i10_i17_and_i19_skipped(tmp_path: Path) -> None:
    root, policy = _two_shard_policy(tmp_path)
    update_child(
        root,
        "FDD-PL-004",
        lambda child: child.model_copy(
            update={"obligation_statement": "Fictional obligation FDD-CF-004."}
        ),
    )
    limited = validate_register(root, policy_path=policy, shards=["PL"])
    assert invariant_status("I-10", limited) == InvariantStatus.SKIPPED.value
    assert invariant_status("I-17", limited) == InvariantStatus.SKIPPED.value
    assert invariant_status("I-19", limited) == InvariantStatus.SKIPPED.value
    viewed = validate_register(root, policy_path=policy, view="fictional-pl")
    assert invariant_status("I-17", viewed) == InvariantStatus.SKIPPED.value
    assert invariant_status("I-19", viewed) == InvariantStatus.SKIPPED.value
    full = validate_register(root, policy_path=policy)
    assert invariant_status("I-10", full) == InvariantStatus.FAIL.value
    assert invariant_status("I-19", full) == InvariantStatus.PASS.value
    assert not any(
        item.code == Code.PENDING_CROSS_SHARD.value and item.severity is FindingSeverity.ERROR
        for item in limited.findings
    )


def _marked_child(child: ChildObligation) -> ChildObligation:
    return child.model_copy(
        update={
            "obligation_statement": f"{_MARKER}\tline\n\x01",
            "acceptance_case": "A" * (_CELL_MAX_CHARS + 1),
        }
    )


def _raising_writer(*_args: object, **_kwargs: object) -> None:
    raise RuntimeError(f"{_MARKER} must stay off the console")


def _cli():
    spec = importlib.util.spec_from_file_location(
        "gl02_register_cli_corrections", Path("scripts/gl02_register.py")
    )
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module
