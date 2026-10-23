"""Project-authored production-shaped FDD workbook fixtures.

The fixtures reproduce only the workbook's structural contract.  All cell
content is fictional and deliberately contains no real-example company data.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

SHEET_ORDER = ("P&L", "Cash Flow", "Liabilities", "Assets")
SHEET_ROW_COUNTS = {
    "P&L": 29,
    "Cash Flow": 13,
    "Liabilities": 21,
    "Assets": 16,
}
SHEET_HEADERS = {
    "P&L": ("Term", "Nature", "Profit & Loss line", "Usual Due Diligence Questions"),
    "Cash Flow": ("Term", "Cash Flow line", "Usual Due Diligence Questions"),
    "Liabilities": (
        "Term",
        "Nature",
        "Balance Sheet line",
        "Usual Due Diligence Questions",
    ),
    "Assets": (
        "Term",
        "Nature",
        "Balance Sheet line",
        "Usual Due Diligence Questions",
    ),
}

# These apparent blank question cells are children of real OOXML merged ranges.
# Import must expand the anchor prose while preserving each physical row trace.
LIABILITIES_MERGED_QUESTION_RANGES = ((4, 6), (17, 18))
CASH_FLOW_EMPTY_LINE_ITEM_ROW = 16


def production_fdd_workbook_bytes(
    *,
    row_delta: int = 0,
    header_row: int = 2,
    missing_sheet: str | None = None,
    sheet_order: tuple[str, ...] = SHEET_ORDER,
    extra_sheet: bool = False,
    duplicate_row: bool = False,
    missing_line_item_trace: bool = False,
    source_marker: str = "fixture-v1",
) -> bytes:
    """Build a fictional workbook with the exact production layout contract.

    ``row_delta`` is applied only to Assets and is intentionally restricted to
    -1/0/+1 so tests can construct the reviewed 78/79/80 row boundary.
    """

    if row_delta not in {-1, 0, 1}:
        raise ValueError("row_delta must be -1, 0, or 1")
    if header_row not in {1, 2, 3}:
        raise ValueError("header_row must be 1, 2, or 3")

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "IDIS fictional Task5 fixture"
    workbook.properties.identifier = source_marker

    if set(sheet_order) != set(SHEET_ORDER) or len(sheet_order) != len(SHEET_ORDER):
        raise ValueError("sheet_order must contain each production sheet exactly once")

    for sheet_name in sheet_order:
        if sheet_name == missing_sheet:
            continue
        sheet = workbook.create_sheet(sheet_name)
        headers = SHEET_HEADERS[sheet_name]
        for column, header in enumerate(headers, start=2):
            sheet.cell(row=header_row, column=column, value=header)

        # Production data begins at row 4.  A deliberately displaced header
        # still leaves data at row 4 so header scanning cannot make it valid.
        count = SHEET_ROW_COUNTS[sheet_name]
        if sheet_name == "Assets":
            count += row_delta
        first_payload: tuple[str, ...] | None = None
        for offset in range(count):
            excel_row = 4 + offset
            term = f"fictional-{sheet_name}-term-{offset + 1:02d}"
            nature = f"fictional-{sheet_name}-nature-{offset + 1:02d}"
            line_item = f"fictional-{sheet_name}-line-{offset + 1:02d}"
            question = f"Review fictional {sheet_name} evidence row {offset + 1:02d}."
            if sheet_name == "Cash Flow" and excel_row == CASH_FLOW_EMPTY_LINE_ITEM_ROW:
                line_item = ""
            if missing_line_item_trace and sheet_name == "Assets" and offset == 0:
                line_item = ""

            if sheet_name == "Cash Flow":
                payload = (term, line_item, question)
            else:
                payload = (term, nature, line_item, question)
            if duplicate_row and sheet_name == "P&L" and offset == 1:
                assert first_payload is not None
                payload = first_payload
            if offset == 0:
                first_payload = payload
            for column, value in enumerate(payload, start=2):
                sheet.cell(row=excel_row, column=column, value=value)

        if sheet_name == "Liabilities":
            for first_row, last_row in LIABILITIES_MERGED_QUESTION_RANGES:
                sheet.merge_cells(
                    start_row=first_row,
                    start_column=5,
                    end_row=last_row,
                    end_column=5,
                )

    if extra_sheet:
        workbook.create_sheet("Fictional Extra")

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
