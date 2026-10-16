"""Independent Task-8.1 oracle for the frozen FDD formula-policy bytes.

This module intentionally does not import a Task-8 materializer or generator.
It translates the approved design into a wished-for artifact contract so the
materialization code has to meet an authority independent from itself.
"""

from __future__ import annotations

import hashlib
import json
import unicodedata
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from idis.methodology.ids import generate_methodology_question_id
from idis.methodology.models import MethodologyType

REPO_ROOT = Path(__file__).resolve().parents[2]
POLICY_ARTIFACT = REPO_ROOT / "contracts" / "fdd_formula_policy_v1.json"
POLICY_FIXTURE = (
    REPO_ROOT / "tests" / "slice100_task8" / "fixtures" / "fdd_formula_policy_v1_hash_v1.json"
)
REVIEW_MAP = (
    REPO_ROOT / "docs" / "plans" / "2026-07-25-slice100-task8-fdd-formula-policy-v1-review-map.md"
)
OPERATOR_WORKBOOK = Path(r"C:\Projects\IDIS\IDIS\financial Due Diligence.xlsx")

WORKBOOK_SHA256 = "a3b83addfff02815c9fd44c4a3dcb7f56ce39410fe182faff12acfb845bd4352"
WORKBOOK_BYTE_COUNT = 24_905
WORKBOOK_SAFE_ROW_COORDINATE_DIGEST = (
    "fab5ab40ab300602b5d9ee59eeb9c86451048d02e419054b0675e753a0ab0bc0"
)
WORKBOOK_SAFE_HEADER_DIGEST = "88c4ea73409b768686dd9dce7d05efe3ccb846bdb13850499caddcd55bc3229c"

SHEETS = ("P&L", "Cash Flow", "Liabilities", "Assets")
SHEET_COUNTS = {"P&L": 29, "Cash Flow": 13, "Liabilities": 21, "Assets": 16}
SHEET_COLUMNS: dict[str, dict[str, int | None]] = {
    "P&L": {"term": 2, "nature": 3, "line_item": 4, "question_text": 5},
    "Cash Flow": {"term": 2, "nature": None, "line_item": 3, "question_text": 4},
    "Liabilities": {"term": 2, "nature": 3, "line_item": 4, "question_text": 5},
    "Assets": {"term": 2, "nature": 3, "line_item": 4, "question_text": 5},
}

POLICY_ID = "fdd-formula-policy-v1"
POLICY_SCHEMA = "fdd-formula-policy-v1"
CANONICALIZATION_ID = "idis-task7-canonical-json-sha256-v1"
FORMULA_ENGINE_ID = "fdd-formula-engine-v1"

EXPECTED_TOP_LEVEL_KEYS = (
    "metadata",
    "canonicalization_contract",
    "workbook_authority",
    "workbook_row_registry",
    "methodology_transition",
    "formula_registry",
    "debt_scope_exclusions",
    "completeness_request_registry",
    "direct_signed_convention_registry",
    "incomplete_binding_failure_schema",
    "financial_output_schema",
    "digest_recipe_registry",
    "activation_assertions",
)

FORMULA_ANCHORS = (
    (1, "GROSS_MARGIN", "P&L", 11, "line_item"),
    (2, "FDD_EBITDA_BRIDGE", "P&L", 23, "p_and_l_control"),
    (3, "FDD_EBIT_BRIDGE", "P&L", 27, "p_and_l_control"),
    (4, "FDD_TAX_RECONCILIATION", "P&L", 30, "p_and_l_control"),
    (5, "FDD_NET_INCOME_BRIDGE", "P&L", 32, "p_and_l_control"),
    (6, "FDD_WORKING_CAPITAL_MOVEMENT", "Cash Flow", 8, "cash_flow_control"),
    (7, "FDD_FREE_CASH_FLOW_CASH_BRIDGE", "Cash Flow", 16, "cash_flow_control"),
    (8, "FDD_DEBT_ROLL_FORWARD", "Cash Flow", 11, "debt_control"),
    (9, "FDD_ASSET_TIE_OUT", "Assets", 19, "sheet_control"),
    (10, "FDD_LIABILITY_TIE_OUT", "Liabilities", 24, "sheet_control"),
)

DIRECT_SIGNED_CONVENTIONS = (
    (
        "fdd.reported_gross_margin_ratio",
        "reported gross profit as a positive fraction of revenue",
        "reported gross loss as a negative fraction",
    ),
    ("fdd.reported_ebitda", "reported EBITDA profit", "reported EBITDA loss"),
    (
        "fdd.other_non_cash_adjustment",
        "adjustment added in the EBIT bridge",
        "adjustment subtracted in the EBIT bridge",
    ),
    ("fdd.reported_ebit", "reported EBIT profit", "reported EBIT loss"),
    (
        "fdd.tax_reconciling_item",
        "increases calculated tax expense",
        "decreases calculated tax expense",
    ),
    (
        "fdd.deferred_tax_adjustment",
        "increases calculated tax expense",
        "decreases calculated tax expense",
    ),
    (
        "fdd.reported_tax_component",
        "reported tax expense",
        "reported tax benefit",
    ),
    ("fdd.reported_net_income", "reported net profit", "reported net loss"),
    (
        "fdd.reported_delta_nwc",
        "movement convention: closing operating NWC exceeds opening operating NWC",
        "movement convention: closing operating NWC is below opening operating NWC",
    ),
    (
        "fdd.nonoperating_working_capital_cash_effect",
        "cash inflow/increase",
        "cash outflow/decrease",
    ),
    (
        "fdd.reported_free_cash_flow",
        "positive free-cash-flow generation",
        "free-cash-flow outflow",
    ),
    (
        "fdd.bank_lease_financing_cash_flow",
        "financing cash inflow/draw",
        "financing repayment/outflow",
    ),
    (
        "fdd.other_financing_cash_flow",
        "financing cash inflow",
        "financing cash outflow",
    ),
    (
        "fdd.recapitalization_cash_flow",
        "recapitalization cash inflow",
        "recapitalization cash outflow",
    ),
    (
        "fdd.other_equity_cash_flow",
        "other equity cash inflow",
        "other equity cash outflow",
    ),
    (
        "fdd.reported_net_cash_flow",
        "reported increase in cash",
        "reported decrease in cash",
    ),
    (
        "fdd.debt_fx_other_noncash",
        "noncash change increasing closing debt",
        "noncash change decreasing closing debt",
    ),
)

COMPLETENESS_REQUESTS = (
    (
        "fdd.tax_reconciling_item_set_complete",
        "P&L",
        30,
        "Return Boolean true only when a cited tax reconciliation or tax schedule "
        "explicitly enumerates the complete set of permanent, rate, and discrete "
        "reconciling items for the stated period and ties that set to the reported "
        "current/deferred tax. If no cited document explicitly establishes "
        "completeness, return EVIDENCE_NOT_FOUND. Never infer completeness from "
        "silence, absence, a calculated residual, or model judgment.",
    ),
    (
        "fdd.debt_rollforward_component_set_complete",
        "Cash Flow",
        11,
        "Return Boolean true only when a cited loan or debt schedule explicitly "
        "enumerates all draws, capitalized interest or accrual, FX or other noncash "
        "changes, repayments, conversions, and writeoffs for the stated period and "
        "ties the schedule's opening and closing balances to Liabilities:9. If no "
        "cited document explicitly establishes completeness, return "
        "EVIDENCE_NOT_FOUND. Never infer completeness from silence, absence, "
        "balancing arithmetic, or model judgment.",
    ),
)

INCOMPLETE_TOP_FIELDS = (
    "schema_version",
    "code",
    "formula_policy_id",
    "formula_policy_sha256",
    "formula_type",
    "formula_ordinal",
    "formula_sha256",
    "anchor_sheet",
    "anchor_row",
    "economic_coordinate",
    "issue_count",
    "issues",
)
INCOMPLETE_ISSUE_FIELDS = (
    "issue_ordinal",
    "issue_kind",
    "binding_ordinal",
    "role",
    "predicate",
    "source_sheet",
    "source_row",
    "economic_coordinate",
    "candidate_factual_output_ids",
)

FINANCIAL_TOP_FIELDS = (
    "schema_version",
    "tenant_id",
    "deal_id",
    "run_id",
    "pipeline_contract",
    "financial_methodology",
    "formula_policy",
    "row_count",
    "rows",
    "reconciliation",
)
FINANCIAL_ROW_FIELDS = (
    "row_ordinal",
    "sheet_ordinal",
    "source_sheet",
    "source_row_number",
    "source_anchor_cell",
    "methodology_version_record_id",
    "methodology_version_id",
    "methodology_question_record_id",
    "methodology_question_id",
    "question_payload_hash",
    "question_text",
    "coverage_record_id",
    "dashboard_row_id",
    "dashboard_output_object_id",
    "disposition",
    "disposition_reason_code",
    "result_records",
    "answer_values",
    "task_gaps",
    "root_links",
    "calculations",
    "row_payload_sha256",
    "root_projection_sha256",
)
FINANCIAL_RESULT_FIELDS = (
    "task_record_id",
    "result_id",
    "output_object_id",
    "output_identity_sha256",
    "output_value_sha256",
    "outcome",
    "result_sha256",
    "searched_scope_sha256",
    "provider_provenance_sha256",
)
FINANCIAL_ANSWER_FIELDS = (
    "output_ordinal",
    "factual_output_record_id",
    "factual_output_object_id",
    "output_identity_sha256",
    "output_value_sha256",
    "predicate",
    "typed_value",
    "period",
    "unit",
    "value_sha256",
    "claim_link_id",
    "claim_id",
    "evidence_links",
    "sanad_links",
)
FINANCIAL_TASK_GAP_FIELDS = (
    "task_record_id",
    "result_id",
    "result_output_object_id",
    "result_output_identity_sha256",
    "result_output_value_sha256",
    "code",
    "finding",
    "searched_scope_sha256",
    "provider_provenance_sha256",
)
FINANCIAL_ROOT_LINK_FIELDS = (
    "root_kind",
    "root_ordinal",
    "output_object_id",
    "output_kind",
    "output_identity_sha256",
    "output_value_sha256",
    "root_link_id",
)
FINANCIAL_RECONCILIATION_FIELDS = (
    "answered_row_count",
    "gap_row_count",
    "red_flag_row_count",
    "result_root_count",
    "evidence_root_count",
    "sanad_root_count",
    "calculation_root_count",
    "adverse_root_count",
    "calculation_count",
    "calculated_only_count",
    "in_balance_count",
    "out_of_balance_count",
    "row_set_sha256",
    "root_set_sha256",
    "calculation_set_sha256",
)

HASHED_COMPONENTS = (
    "metadata",
    "canonicalization_contract",
    "workbook_authority",
    "workbook_row_registry",
    "methodology_transition",
    "formula_registry",
    "debt_scope_exclusions",
    "completeness_request_registry",
    "direct_signed_convention_registry",
    "incomplete_binding_failure_schema",
    "financial_output_schema",
    "activation_assertions",
)


class OracleError(AssertionError):
    """Safe failure carrying only a stable contract code."""


def _safe_require(condition: bool, code: str) -> None:
    if not condition:
        raise OracleError(code)


def canonical_json(value: Any) -> bytes:
    """Render the frozen Task-7 canonical JSON subset used by the policy."""

    _validate_canonical_value(value)
    return json.dumps(
        value,
        ensure_ascii=False,
        allow_nan=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def _validate_canonical_value(value: Any) -> None:
    if isinstance(value, float):
        raise OracleError("TASK8_POLICY_FLOAT_FORBIDDEN")
    if isinstance(value, str):
        _safe_require(
            unicodedata.normalize("NFC", value) == value,
            "TASK8_POLICY_UNICODE_NOT_NFC",
        )
        return
    if value is None or isinstance(value, (bool, int)):
        return
    if isinstance(value, list):
        for item in value:
            _validate_canonical_value(item)
        return
    if isinstance(value, dict):
        for key, item in value.items():
            _safe_require(isinstance(key, str), "TASK8_POLICY_NON_STRING_KEY")
            _validate_canonical_value(key)
            _validate_canonical_value(item)
        return
    raise OracleError("TASK8_POLICY_VALUE_TYPE_INVALID")


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_canonical(value: Any) -> str:
    return sha256_bytes(canonical_json(value))


def load_json_safe(path: Path, missing_code: str) -> dict[str, Any]:
    _safe_require(path.is_file(), f"{missing_code}:{path.name}")
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise OracleError(f"{missing_code}_INVALID") from exc
    _safe_require(isinstance(value, dict), f"{missing_code}_INVALID")
    return value


def load_policy() -> dict[str, Any]:
    return load_json_safe(POLICY_ARTIFACT, "TASK8_FDD_POLICY_ARTIFACT_NOT_IMPLEMENTED")


def load_fixture() -> dict[str, Any]:
    return load_json_safe(POLICY_FIXTURE, "TASK8_FDD_POLICY_FIXTURE_NOT_IMPLEMENTED")


def require_review_map() -> bytes:
    _safe_require(
        REVIEW_MAP.is_file(),
        f"TASK8_FDD_REVIEW_MAP_NOT_IMPLEMENTED:{REVIEW_MAP.name}",
    )
    return REVIEW_MAP.read_bytes()


def validate_top_level(policy: Mapping[str, Any]) -> None:
    _safe_require(
        tuple(policy) == EXPECTED_TOP_LEVEL_KEYS,
        "TASK8_FDD_POLICY_TOP_LEVEL_INVALID",
    )
    _safe_require(
        policy["metadata"] == expected_metadata(),
        "TASK8_FDD_POLICY_METADATA_INVALID",
    )


def expected_metadata() -> dict[str, Any]:
    return {
        "artifact_schema": POLICY_SCHEMA,
        "policy_id": POLICY_ID,
        "policy_version": 1,
        "append_frozen": True,
        "activation_status": "INACTIVE",
        "canonicalization_contract_id": CANONICALIZATION_ID,
        "formula_engine_id": FORMULA_ENGINE_ID,
        "workbook_sha256": WORKBOOK_SHA256,
        "workbook_byte_count": WORKBOOK_BYTE_COUNT,
        "workbook_safe_row_coordinate_digest": WORKBOOK_SAFE_ROW_COORDINATE_DIGEST,
        "workbook_safe_header_digest": WORKBOOK_SAFE_HEADER_DIGEST,
        "question_count": 79,
        "required_question_count": 10,
        "empty_requirement_count": 69,
        "formula_count": 10,
        "expanded_source_binding_count": 107,
        "distinct_source_mapping_count": 101,
        "parent_binding_count": 9,
        "direct_signed_binding_count": 19,
        "completeness_binding_count": 2,
    }


def expected_canonicalization_contract() -> dict[str, Any]:
    return {
        "contract_id": CANONICALIZATION_ID,
        "encoding": "UTF-8",
        "bom": "FORBIDDEN",
        "unicode": "NFC",
        "object_key_order": "UNICODE_CODEPOINT_LEXICAL",
        "separators": [",", ":"],
        "insignificant_whitespace": "FORBIDDEN",
        "json_float": "FORBIDDEN",
        "ordered_arrays": "PRESERVED",
        "decimal": {
            "representation": "FINITE_NON_EXPONENT_BASE10_STRING",
            "trailing_fractional_zeroes": "REMOVED",
            "negative_zero": "0",
        },
        "uuid": "LOWERCASE_CANONICAL",
        "sha256": "LOWERCASE_HEX_64",
        "timestamp": "UTC_RFC3339_SIX_FRACTIONAL_DIGITS_Z",
    }


def read_operator_rows() -> list[dict[str, Any]]:
    """Read the pinned workbook without using the production FDD importer."""

    _safe_require(OPERATOR_WORKBOOK.is_file(), "TASK8_OPERATOR_WORKBOOK_MISSING")
    raw = OPERATOR_WORKBOOK.read_bytes()
    _safe_require(len(raw) == WORKBOOK_BYTE_COUNT, "TASK8_OPERATOR_WORKBOOK_SIZE_DRIFT")
    _safe_require(
        sha256_bytes(raw) == WORKBOOK_SHA256,
        "TASK8_OPERATOR_WORKBOOK_HASH_DRIFT",
    )
    workbook = load_workbook(OPERATOR_WORKBOOK, data_only=False, read_only=False)
    try:
        _safe_require(tuple(workbook.sheetnames) == SHEETS, "TASK8_WORKBOOK_SHEET_DRIFT")
        rows: list[dict[str, Any]] = []
        ordinal = 0
        formula_count = 0
        for sheet_ordinal, sheet_name in enumerate(SHEETS, start=1):
            sheet = workbook[sheet_name]
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula_count += 1
            sheet_rows = _read_sheet_rows(sheet_name, sheet, sheet_ordinal, ordinal)
            rows.extend(sheet_rows)
            ordinal += len(sheet_rows)
        _safe_require(formula_count == 0, "TASK8_WORKBOOK_FORMULA_CELL_DRIFT")
        _safe_require(len(rows) == 79, "TASK8_WORKBOOK_ROW_COUNT_DRIFT")
        return rows
    finally:
        workbook.close()


def _read_sheet_rows(
    sheet_name: str,
    sheet: Worksheet,
    sheet_ordinal: int,
    prior_count: int,
) -> list[dict[str, Any]]:
    columns = SHEET_COLUMNS[sheet_name]
    rows: list[dict[str, Any]] = []
    for source_row in range(4, sheet.max_row + 1):
        term, _term_anchor = _resolved_cell(sheet, source_row, columns["term"])
        nature, _nature_anchor = _resolved_cell(sheet, source_row, columns["nature"])
        line_item, _line_anchor = _resolved_cell(sheet, source_row, columns["line_item"])
        question_text, source_anchor_cell = _resolved_cell(
            sheet,
            source_row,
            columns["question_text"],
        )
        if not any((term, nature, line_item, question_text)):
            continue
        methodology_question_id = generate_methodology_question_id(
            methodology_type=MethodologyType.FINANCIAL_DD,
            section=sheet_name,
            sheet_or_section=sheet_name,
            row_number=source_row,
            line_item=line_item,
            question_text=question_text,
            source_hash=WORKBOOK_SHA256,
        )
        row_ordinal = prior_count + len(rows) + 1
        required = [
            {
                "formula_ordinal": ordinal,
                "formula_type": formula_type,
                "required": True,
                "scope": scope,
            }
            for ordinal, formula_type, anchor_sheet, anchor_row, scope in FORMULA_ANCHORS
            if (anchor_sheet, anchor_row) == (sheet_name, source_row)
        ]
        row = {
            "row_ordinal": row_ordinal,
            "sheet_ordinal": sheet_ordinal,
            "source_sheet": sheet_name,
            "source_row": source_row,
            "source_anchor_cell": source_anchor_cell,
            "term": term,
            "nature": nature or None,
            "line_item": line_item or None,
            "question_text": question_text,
            "methodology_question_id": methodology_question_id,
            "required_calculations": required,
        }
        row["source_row_sha256"] = source_row_sha256(row)
        rows.append(row)
    _safe_require(
        len(rows) == SHEET_COUNTS[sheet_name],
        "TASK8_WORKBOOK_SHEET_ROW_COUNT_DRIFT",
    )
    return rows


def _resolved_cell(
    sheet: Worksheet,
    source_row: int,
    column: int | None,
) -> tuple[str, str]:
    if column is None:
        return "", ""
    anchor_row = source_row
    anchor_column = column
    value = sheet.cell(row=source_row, column=column).value
    if value is None:
        for merged in sheet.merged_cells.ranges:
            if (
                merged.min_row <= source_row <= merged.max_row
                and merged.min_col <= column <= merged.max_col
            ):
                anchor_row = merged.min_row
                anchor_column = merged.min_col
                value = sheet.cell(row=anchor_row, column=anchor_column).value
                break
    return (
        "" if value is None else str(value).strip(),
        f"{get_column_letter(anchor_column)}{anchor_row}",
    )


def source_row_sha256(row: Mapping[str, Any]) -> str:
    return sha256_canonical(
        [
            "fdd-workbook-row-v1",
            row["row_ordinal"],
            row["sheet_ordinal"],
            row["source_sheet"],
            row["source_row"],
            row["source_anchor_cell"],
            row["term"],
            row["nature"],
            row["line_item"],
            row["question_text"],
            row["methodology_question_id"],
        ]
    )


def expected_workbook_authority() -> dict[str, Any]:
    return {
        "schema_version": "fdd-workbook-authority-v1",
        "source_name": "configured-fdd-workbook",
        "source_type": "operator_workbook",
        "byte_count": WORKBOOK_BYTE_COUNT,
        "sha256": WORKBOOK_SHA256,
        "safe_row_coordinate_digest": WORKBOOK_SAFE_ROW_COORDINATE_DIGEST,
        "safe_header_digest": WORKBOOK_SAFE_HEADER_DIGEST,
        "header_row": 2,
        "first_data_row": 4,
        "formula_cell_count": 0,
        "sheet_registry": [
            {
                "sheet_ordinal": ordinal,
                "sheet_name": name,
                "question_count": SHEET_COUNTS[name],
            }
            for ordinal, name in enumerate(SHEETS, start=1)
        ],
    }


def _source(
    role: str,
    sheet: str,
    rows: int | Iterable[int],
    predicate: str,
    *,
    value_schema: str = "FddMonetaryValueV1",
    sign: str = "POSITIVE_MAGNITUDE",
    cardinality: str = "EXACTLY_ONE",
    period_role: str = "FLOW_RANGE",
    unit_rule: str = "COMMON_MONETARY_UNIT",
) -> list[dict[str, Any]]:
    row_tuple = (rows,) if isinstance(rows, int) else tuple(rows)
    return [
        {
            "binding_kind": "SOURCE",
            "role": role,
            "predicate": predicate,
            "source_sheet": sheet,
            "source_row": row,
            "value_schema": value_schema,
            "sign_transform": sign,
            "cardinality": cardinality,
            "period_role": period_role,
            "currency_rule": (
                None
                if value_schema in {"FddSignedRatioValueV1", "FddBooleanValueV1"}
                else "COMMON_ISO4217_AND_SCALE"
            ),
            "unit_rule": unit_rule,
            "direct_signed_convention_id": predicate if sign == "DIRECT_SIGNED" else None,
            "completeness_request_id": (predicate if value_schema == "FddBooleanValueV1" else None),
        }
        for row in row_tuple
    ]


def _parent(
    role: str,
    formula_type: str,
    *,
    coordinate_rule: str = "EXACT_SAME_COORDINATE",
) -> dict[str, Any]:
    return {
        "binding_kind": "PARENT_CALCULATION",
        "role": role,
        "parent_formula_type": formula_type,
        "coordinate_rule": coordinate_rule,
        "cardinality": "EXACTLY_ONE",
    }


def _with_ordinals(bindings: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    return [
        {"binding_ordinal": ordinal, **dict(binding)}
        for ordinal, binding in enumerate(bindings, start=1)
    ]


def expected_formula_rows() -> list[dict[str, Any]]:
    """Return the ten exact formula rows authorized by the design ruling."""

    formulas: list[dict[str, Any]] = []

    gross_bindings = [
        *_source("recurring_revenue", "P&L", 4, "fdd.recurring_revenue"),
        *_source("nonrecurring_revenue", "P&L", 5, "fdd.nonrecurring_revenue"),
        *_source(
            "cost_of_sales_components",
            "P&L",
            (7, 8, 9, 10),
            "fdd.cost_of_sales_component",
            cardinality="ORDERED_COMPONENT_SET",
        ),
        *_source(
            "reported_gross_margin_ratio",
            "P&L",
            11,
            "fdd.reported_gross_margin_ratio",
            value_schema="FddSignedRatioValueV1",
            sign="DIRECT_SIGNED",
            unit_rule="RATIO_FRACTION",
        ),
    ]
    formulas.append(
        _formula(
            1,
            "GROSS_MARGIN",
            "gross-margin-v1",
            "P&L",
            11,
            "line_item",
            gross_bindings,
            [],
            [
                [
                    "assign",
                    "operating_revenue",
                    [
                        "add",
                        ["input", "recurring_revenue"],
                        ["input", "nonrecurring_revenue"],
                    ],
                ],
                [
                    "assign",
                    "cost_of_sales",
                    ["sum", "cost_of_sales_components"],
                ],
                [
                    "assign",
                    "gross_profit",
                    [
                        "subtract",
                        ["value", "operating_revenue"],
                        ["value", "cost_of_sales"],
                    ],
                ],
                [
                    "assign",
                    "gross_margin_ratio",
                    [
                        "divide",
                        ["value", "gross_profit"],
                        ["value", "operating_revenue"],
                    ],
                ],
                [
                    "assign",
                    "gross_margin_ratio_variance",
                    [
                        "subtract",
                        ["value", "gross_margin_ratio"],
                        ["input", "reported_gross_margin_ratio"],
                    ],
                ],
            ],
            denominator_rule={
                "name": "operating_revenue",
                "constraint": "STRICTLY_POSITIVE",
            },
            outputs=[
                _output("gross_margin_ratio", "ratio", 4, primary=True),
                _output("gross_profit", "money", 2),
                _output("operating_revenue", "money", 2),
                _output("cost_of_sales", "money", 2),
                _output("reported_gross_margin_ratio", "ratio", 4),
                _output("gross_margin_ratio_variance", "ratio", 4, comparison=True),
            ],
        )
    )

    ebitda_bindings = [
        *_source(
            "operating_income_components",
            "P&L",
            (4, 5, 6),
            "fdd.operating_income_component",
            cardinality="ORDERED_COMPONENT_SET",
        ),
        *_source(
            "cost_of_sales_components",
            "P&L",
            (7, 8, 9, 10),
            "fdd.cost_of_sales_component",
            cardinality="ORDERED_COMPONENT_SET",
        ),
        *_source(
            "operating_expense_components",
            "P&L",
            range(12, 23),
            "fdd.operating_expense_component",
            cardinality="ORDERED_COMPONENT_SET",
        ),
        _parent("operating_revenue", "GROSS_MARGIN"),
        *_source(
            "reported_ebitda",
            "P&L",
            23,
            "fdd.reported_ebitda",
            sign="DIRECT_SIGNED",
        ),
    ]
    formulas.append(
        _formula(
            2,
            "FDD_EBITDA_BRIDGE",
            "fdd-ebitda-bridge-v1",
            "P&L",
            23,
            "p_and_l_control",
            ebitda_bindings,
            ["GROSS_MARGIN"],
            [
                [
                    "assign",
                    "calculated_ebitda",
                    [
                        "subtract",
                        [
                            "subtract",
                            ["sum", "operating_income_components"],
                            ["sum", "cost_of_sales_components"],
                        ],
                        ["sum", "operating_expense_components"],
                    ],
                ],
                [
                    "assign",
                    "ebitda_margin_ratio",
                    [
                        "divide",
                        ["value", "calculated_ebitda"],
                        ["parent", "operating_revenue"],
                    ],
                ],
                [
                    "assign",
                    "ebitda_variance",
                    [
                        "subtract",
                        ["value", "calculated_ebitda"],
                        ["input", "reported_ebitda"],
                    ],
                ],
            ],
            denominator_rule={
                "name": "operating_revenue",
                "constraint": "STRICTLY_POSITIVE",
            },
            outputs=[
                _output("calculated_ebitda", "money", 2, primary=True),
                _output("reported_ebitda", "money", 2),
                _output("ebitda_margin_ratio", "ratio", 4),
                _output("ebitda_variance", "money", 2, comparison=True),
            ],
        )
    )

    ebit_bindings = [
        _parent("calculated_ebitda", "FDD_EBITDA_BRIDGE"),
        _parent("operating_revenue", "GROSS_MARGIN"),
        *_source("depreciation", "P&L", 24, "fdd.depreciation"),
        *_source(
            "amortization_and_impairment",
            "P&L",
            25,
            "fdd.amortization_impairment",
        ),
        *_source(
            "other_non_cash_adjustment",
            "P&L",
            26,
            "fdd.other_non_cash_adjustment",
            sign="DIRECT_SIGNED",
        ),
        *_source(
            "reported_ebit",
            "P&L",
            27,
            "fdd.reported_ebit",
            sign="DIRECT_SIGNED",
        ),
    ]
    formulas.append(
        _formula(
            3,
            "FDD_EBIT_BRIDGE",
            "fdd-ebit-bridge-v1",
            "P&L",
            27,
            "p_and_l_control",
            ebit_bindings,
            ["FDD_EBITDA_BRIDGE", "GROSS_MARGIN"],
            [
                [
                    "assign",
                    "calculated_ebit",
                    [
                        "add",
                        [
                            "subtract",
                            [
                                "subtract",
                                ["parent", "calculated_ebitda"],
                                ["input", "depreciation"],
                            ],
                            ["input", "amortization_and_impairment"],
                        ],
                        ["input", "other_non_cash_adjustment"],
                    ],
                ],
                [
                    "assign",
                    "ebit_margin_ratio",
                    [
                        "divide",
                        ["value", "calculated_ebit"],
                        ["parent", "operating_revenue"],
                    ],
                ],
                [
                    "assign",
                    "ebit_variance",
                    [
                        "subtract",
                        ["value", "calculated_ebit"],
                        ["input", "reported_ebit"],
                    ],
                ],
            ],
            denominator_rule={
                "name": "operating_revenue",
                "constraint": "STRICTLY_POSITIVE",
            },
            outputs=[
                _output("calculated_ebit", "money", 2, primary=True),
                _output("reported_ebit", "money", 2),
                _output("ebit_margin_ratio", "ratio", 4),
                _output("ebit_variance", "money", 2, comparison=True),
            ],
        )
    )

    tax_bindings = [
        _parent("calculated_ebit", "FDD_EBIT_BRIDGE"),
        *_source(
            "interest_expense_components",
            "P&L",
            (28, 29),
            "fdd.interest_expense_component",
            cardinality="ORDERED_COMPONENT_SET",
        ),
        *_source(
            "statutory_tax_rate_ratio",
            "P&L",
            30,
            "fdd.statutory_tax_rate_ratio",
            value_schema="FddRatioValueV1",
            unit_rule="RATIO_FRACTION",
        ),
        *_source(
            "tax_reconciling_items",
            "P&L",
            30,
            "fdd.tax_reconciling_item",
            sign="DIRECT_SIGNED",
            cardinality="ORDERED_ZERO_OR_MORE",
        ),
        *_source(
            "reconciling_set_complete",
            "P&L",
            30,
            "fdd.tax_reconciling_item_set_complete",
            value_schema="FddBooleanValueV1",
            cardinality="EXACTLY_ONE_TRUE",
            unit_rule="BOOLEAN",
        ),
        *_source(
            "deferred_tax_adjustment",
            "P&L",
            31,
            "fdd.deferred_tax_adjustment",
            sign="DIRECT_SIGNED",
        ),
        *_source(
            "reported_tax_components",
            "P&L",
            (30, 31),
            "fdd.reported_tax_component",
            sign="DIRECT_SIGNED",
            cardinality="ORDERED_COMPONENT_SET",
        ),
    ]
    formulas.append(
        _formula(
            4,
            "FDD_TAX_RECONCILIATION",
            "fdd-tax-reconciliation-v1",
            "P&L",
            30,
            "p_and_l_control",
            tax_bindings,
            ["FDD_EBIT_BRIDGE"],
            [
                [
                    "assign",
                    "profit_before_tax",
                    [
                        "subtract",
                        ["parent", "calculated_ebit"],
                        ["sum", "interest_expense_components"],
                    ],
                ],
                [
                    "assign",
                    "calculated_tax",
                    [
                        "add",
                        [
                            "add",
                            [
                                "multiply",
                                ["value", "profit_before_tax"],
                                ["input", "statutory_tax_rate_ratio"],
                            ],
                            ["sum", "tax_reconciling_items"],
                        ],
                        ["input", "deferred_tax_adjustment"],
                    ],
                ],
                [
                    "assign",
                    "reported_total_tax",
                    ["sum", "reported_tax_components"],
                ],
                [
                    "assign",
                    "tax_variance",
                    [
                        "subtract",
                        ["value", "calculated_tax"],
                        ["value", "reported_total_tax"],
                    ],
                ],
            ],
            denominator_rule={
                "name": "statutory_tax_rate_ratio",
                "constraint": "CLOSED_INTERVAL_0_1",
            },
            outputs=[
                _output("calculated_tax", "money", 2, primary=True),
                _output("profit_before_tax", "money", 2),
                _output("reported_total_tax", "money", 2),
                _output("tax_variance", "money", 2, comparison=True),
            ],
        )
    )

    net_bindings = [
        _parent("calculated_ebit", "FDD_EBIT_BRIDGE"),
        _parent("calculated_tax", "FDD_TAX_RECONCILIATION"),
        _parent("operating_revenue", "GROSS_MARGIN"),
        *_source(
            "interest_expense_components",
            "P&L",
            (28, 29),
            "fdd.interest_expense_component",
            cardinality="ORDERED_COMPONENT_SET",
        ),
        *_source(
            "reported_net_income",
            "P&L",
            32,
            "fdd.reported_net_income",
            sign="DIRECT_SIGNED",
        ),
    ]
    formulas.append(
        _formula(
            5,
            "FDD_NET_INCOME_BRIDGE",
            "fdd-net-income-bridge-v1",
            "P&L",
            32,
            "p_and_l_control",
            net_bindings,
            ["FDD_EBIT_BRIDGE", "FDD_TAX_RECONCILIATION", "GROSS_MARGIN"],
            [
                [
                    "assign",
                    "calculated_net_income",
                    [
                        "subtract",
                        [
                            "subtract",
                            ["parent", "calculated_ebit"],
                            ["sum", "interest_expense_components"],
                        ],
                        ["parent", "calculated_tax"],
                    ],
                ],
                [
                    "assign",
                    "net_income_margin_ratio",
                    [
                        "divide",
                        ["value", "calculated_net_income"],
                        ["parent", "operating_revenue"],
                    ],
                ],
                [
                    "assign",
                    "net_income_variance",
                    [
                        "subtract",
                        ["value", "calculated_net_income"],
                        ["input", "reported_net_income"],
                    ],
                ],
            ],
            denominator_rule={
                "name": "operating_revenue",
                "constraint": "STRICTLY_POSITIVE",
            },
            outputs=[
                _output("calculated_net_income", "money", 2, primary=True),
                _output("reported_net_income", "money", 2),
                _output("net_income_margin_ratio", "ratio", 4),
                _output("net_income_variance", "money", 2, comparison=True),
            ],
        )
    )

    nwc_bindings = [
        *_source(
            "operating_current_assets",
            "Assets",
            (13, 14, 15, 16),
            "fdd.operating_current_asset_balance",
            cardinality="EXACTLY_TWO_OPENING_CLOSING",
            period_role="OPENING_AND_CLOSING_AS_OF",
        ),
        *_source(
            "operating_current_liabilities",
            "Liabilities",
            (12, 13, 14, 15, 16),
            "fdd.operating_current_liability_balance",
            cardinality="EXACTLY_TWO_OPENING_CLOSING",
            period_role="OPENING_AND_CLOSING_AS_OF",
        ),
        *_source(
            "reported_delta_nwc",
            "Cash Flow",
            8,
            "fdd.reported_delta_nwc",
            sign="DIRECT_SIGNED",
        ),
    ]
    formulas.append(
        _formula(
            6,
            "FDD_WORKING_CAPITAL_MOVEMENT",
            "fdd-working-capital-movement-v1",
            "Cash Flow",
            8,
            "cash_flow_control",
            nwc_bindings,
            [],
            [
                [
                    "assign",
                    "opening_operating_working_capital",
                    [
                        "subtract",
                        ["sum_opening", "operating_current_assets"],
                        ["sum_opening", "operating_current_liabilities"],
                    ],
                ],
                [
                    "assign",
                    "closing_operating_working_capital",
                    [
                        "subtract",
                        ["sum_closing", "operating_current_assets"],
                        ["sum_closing", "operating_current_liabilities"],
                    ],
                ],
                [
                    "assign",
                    "calculated_delta_nwc",
                    [
                        "subtract",
                        ["value", "closing_operating_working_capital"],
                        ["value", "opening_operating_working_capital"],
                    ],
                ],
                [
                    "assign",
                    "nwc_cash_effect",
                    ["negate", ["value", "calculated_delta_nwc"]],
                ],
                [
                    "assign",
                    "nwc_variance",
                    [
                        "subtract",
                        ["value", "calculated_delta_nwc"],
                        ["input", "reported_delta_nwc"],
                    ],
                ],
            ],
            period_rule="EXACT_OPENING_CLOSING_AS_OF_AND_MATCHING_FLOW_RANGE",
            outputs=[
                _output("calculated_delta_nwc", "money", 2, primary=True),
                _output("opening_operating_working_capital", "money", 2),
                _output("closing_operating_working_capital", "money", 2),
                _output("nwc_cash_effect", "money", 2),
                _output("reported_delta_nwc", "money", 2),
                _output("nwc_variance", "money", 2, comparison=True),
            ],
        )
    )

    cash_bindings = [
        *_source("cash_ebitda", "Cash Flow", 4, "fdd.cash_ebitda"),
        *_source(
            "cash_tax_and_nonrecurring",
            "Cash Flow",
            5,
            "fdd.cash_tax_and_nonrecurring",
        ),
        *_source("tangible_capex", "Cash Flow", 6, "fdd.tangible_capex"),
        *_source("intangible_capex", "Cash Flow", 7, "fdd.intangible_capex"),
        _parent("nwc_cash_effect", "FDD_WORKING_CAPITAL_MOVEMENT"),
        *_source(
            "nonoperating_working_capital_cash_effect",
            "Cash Flow",
            9,
            "fdd.nonoperating_working_capital_cash_effect",
            sign="DIRECT_SIGNED",
        ),
        *_source(
            "reported_free_cash_flow",
            "Cash Flow",
            10,
            "fdd.reported_free_cash_flow",
            sign="DIRECT_SIGNED",
        ),
        *_source(
            "bank_lease_financing_cash_flow",
            "Cash Flow",
            11,
            "fdd.bank_lease_financing_cash_flow",
            sign="DIRECT_SIGNED",
        ),
        *_source(
            "other_financing_cash_flow",
            "Cash Flow",
            12,
            "fdd.other_financing_cash_flow",
            sign="DIRECT_SIGNED",
        ),
        *_source(
            "recapitalization_cash_flow",
            "Cash Flow",
            13,
            "fdd.recapitalization_cash_flow",
            sign="DIRECT_SIGNED",
        ),
        *_source("dividends_paid", "Cash Flow", 14, "fdd.dividends_paid"),
        *_source(
            "other_equity_cash_flow",
            "Cash Flow",
            15,
            "fdd.other_equity_cash_flow",
            sign="DIRECT_SIGNED",
        ),
        *_source(
            "reported_net_cash_flow",
            "Cash Flow",
            16,
            "fdd.reported_net_cash_flow",
            sign="DIRECT_SIGNED",
        ),
        *_source(
            "unrestricted_cash_balance",
            "Assets",
            19,
            "fdd.unrestricted_cash_balance",
            cardinality="EXACTLY_TWO_OPENING_CLOSING",
            period_role="OPENING_AND_CLOSING_AS_OF",
        ),
    ]
    formulas.append(
        _formula(
            7,
            "FDD_FREE_CASH_FLOW_CASH_BRIDGE",
            "fdd-free-cash-flow-cash-bridge-v1",
            "Cash Flow",
            16,
            "cash_flow_control",
            cash_bindings,
            ["FDD_WORKING_CAPITAL_MOVEMENT"],
            [
                [
                    "assign",
                    "calculated_fcf",
                    [
                        "add",
                        [
                            "subtract",
                            [
                                "subtract",
                                [
                                    "subtract",
                                    ["input", "cash_ebitda"],
                                    ["input", "cash_tax_and_nonrecurring"],
                                ],
                                ["input", "tangible_capex"],
                            ],
                            ["input", "intangible_capex"],
                        ],
                        [
                            "add",
                            ["parent", "nwc_cash_effect"],
                            ["input", "nonoperating_working_capital_cash_effect"],
                        ],
                    ],
                ],
                [
                    "assign",
                    "calculated_net_cash_flow",
                    [
                        "add",
                        [
                            "subtract",
                            [
                                "add",
                                [
                                    "add",
                                    [
                                        "add",
                                        ["value", "calculated_fcf"],
                                        ["input", "bank_lease_financing_cash_flow"],
                                    ],
                                    ["input", "other_financing_cash_flow"],
                                ],
                                ["input", "recapitalization_cash_flow"],
                            ],
                            ["input", "dividends_paid"],
                        ],
                        ["input", "other_equity_cash_flow"],
                    ],
                ],
                [
                    "assign",
                    "calculated_closing_cash",
                    [
                        "add",
                        ["input_opening", "unrestricted_cash_balance"],
                        ["value", "calculated_net_cash_flow"],
                    ],
                ],
                [
                    "assign",
                    "free_cash_flow_variance",
                    [
                        "subtract",
                        ["value", "calculated_fcf"],
                        ["input", "reported_free_cash_flow"],
                    ],
                ],
                [
                    "assign",
                    "net_cash_flow_variance",
                    [
                        "subtract",
                        ["value", "calculated_net_cash_flow"],
                        ["input", "reported_net_cash_flow"],
                    ],
                ],
                [
                    "assign",
                    "closing_cash_variance",
                    [
                        "subtract",
                        ["value", "calculated_closing_cash"],
                        ["input_closing", "unrestricted_cash_balance"],
                    ],
                ],
            ],
            period_rule="EXACT_OPENING_CLOSING_AS_OF_AND_MATCHING_FLOW_RANGE",
            outputs=[
                _output("calculated_fcf", "money", 2, primary=True),
                _output("calculated_net_cash_flow", "money", 2),
                _output("calculated_closing_cash", "money", 2),
                _output("reported_free_cash_flow", "money", 2),
                _output("reported_net_cash_flow", "money", 2),
                _output("reported_closing_cash", "money", 2),
                _output("free_cash_flow_variance", "money", 2, comparison=True),
                _output("net_cash_flow_variance", "money", 2, comparison=True),
                _output("closing_cash_variance", "money", 2, comparison=True),
            ],
        )
    )

    debt_bindings = [
        *_source(
            "interest_bearing_debt_balance",
            "Liabilities",
            9,
            "fdd.interest_bearing_debt_balance",
            cardinality="EXACTLY_TWO_OPENING_CLOSING",
            period_role="OPENING_AND_CLOSING_AS_OF",
        ),
        *_source(
            "draws",
            "Cash Flow",
            (11, 12),
            "fdd.debt_draw",
            cardinality="ORDERED_ZERO_OR_MORE",
        ),
        *_source(
            "capitalized_interest",
            "Cash Flow",
            (11, 12),
            "fdd.debt_capitalized_interest",
            cardinality="ORDERED_ZERO_OR_MORE",
        ),
        *_source(
            "fx_and_other_noncash_changes",
            "Cash Flow",
            (11, 12),
            "fdd.debt_fx_other_noncash",
            sign="DIRECT_SIGNED",
            cardinality="ORDERED_ZERO_OR_MORE",
        ),
        *_source(
            "repayments",
            "Cash Flow",
            (11, 12),
            "fdd.debt_repayment",
            cardinality="ORDERED_ZERO_OR_MORE",
        ),
        *_source(
            "conversions_and_writeoffs",
            "Cash Flow",
            (11, 12),
            "fdd.debt_conversion_writeoff",
            cardinality="ORDERED_ZERO_OR_MORE",
        ),
        *_source(
            "movement_set_complete",
            "Cash Flow",
            11,
            "fdd.debt_rollforward_component_set_complete",
            value_schema="FddBooleanValueV1",
            cardinality="EXACTLY_ONE_TRUE",
            unit_rule="BOOLEAN",
        ),
    ]
    formulas.append(
        _formula(
            8,
            "FDD_DEBT_ROLL_FORWARD",
            "fdd-debt-roll-forward-v1",
            "Cash Flow",
            11,
            "debt_control",
            debt_bindings,
            [],
            [
                [
                    "assign",
                    "calculated_closing_debt",
                    [
                        "subtract",
                        [
                            "subtract",
                            [
                                "add",
                                [
                                    "add",
                                    [
                                        "add",
                                        ["input_opening", "interest_bearing_debt_balance"],
                                        ["sum", "draws"],
                                    ],
                                    ["sum", "capitalized_interest"],
                                ],
                                ["sum", "fx_and_other_noncash_changes"],
                            ],
                            ["sum", "repayments"],
                        ],
                        ["sum", "conversions_and_writeoffs"],
                    ],
                ],
                [
                    "assign",
                    "debt_variance",
                    [
                        "subtract",
                        ["value", "calculated_closing_debt"],
                        ["input_closing", "interest_bearing_debt_balance"],
                    ],
                ],
            ],
            period_rule="EXACT_OPENING_CLOSING_AS_OF_AND_MATCHING_FLOW_RANGE",
            outputs=[
                _output("calculated_closing_debt", "money", 2, primary=True),
                _output("reported_closing_debt", "money", 2),
                _output("debt_variance", "money", 2, comparison=True),
            ],
        )
    )

    asset_bindings = _source(
        "asset_component_balances",
        "Assets",
        range(4, 20),
        "fdd.asset_component_balance",
        cardinality="ORDERED_COMPONENT_SET",
        period_role="CLOSING_AS_OF",
    )
    formulas.append(
        _formula(
            9,
            "FDD_ASSET_TIE_OUT",
            "fdd-asset-tie-out-v1",
            "Assets",
            19,
            "sheet_control",
            asset_bindings,
            [],
            [
                [
                    "assign",
                    "calculated_total_assets",
                    ["sum", "asset_component_balances"],
                ]
            ],
            period_rule="EXACT_CLOSING_AS_OF",
            outputs=[
                _output("calculated_total_assets", "money", 2, primary=True),
            ],
            disposition_rule="CALCULATED_ONLY",
        )
    )

    liability_bindings = [
        *_source(
            "liability_equity_component_balances",
            "Liabilities",
            range(4, 19),
            "fdd.liability_equity_component_balance",
            cardinality="ORDERED_COMPONENT_SET",
            period_role="CLOSING_AS_OF",
        ),
        _parent(
            "calculated_total_assets",
            "FDD_ASSET_TIE_OUT",
            coordinate_rule="EXACT_SAME_AS_OF_DATE",
        ),
    ]
    formulas.append(
        _formula(
            10,
            "FDD_LIABILITY_TIE_OUT",
            "fdd-liability-tie-out-v1",
            "Liabilities",
            24,
            "sheet_control",
            liability_bindings,
            ["FDD_ASSET_TIE_OUT"],
            [
                [
                    "assign",
                    "calculated_total_liabilities_and_equity",
                    ["sum", "liability_equity_component_balances"],
                ],
                [
                    "assign",
                    "balance_sheet_variance",
                    [
                        "subtract",
                        ["parent", "calculated_total_assets"],
                        ["value", "calculated_total_liabilities_and_equity"],
                    ],
                ],
            ],
            period_rule="EXACT_CLOSING_AS_OF",
            outputs=[
                _output(
                    "calculated_total_liabilities_and_equity",
                    "money",
                    2,
                    primary=True,
                ),
                _output("calculated_total_assets", "money", 2),
                _output("balance_sheet_variance", "money", 2, comparison=True),
            ],
        )
    )
    return formulas


def _output(
    name: str,
    unit: str,
    scale: int,
    *,
    primary: bool = False,
    comparison: bool = False,
) -> dict[str, Any]:
    return {
        "name": name,
        "unit": unit,
        "currency": None if unit == "ratio" else "COORDINATE_BASE_CURRENCY",
        "scale": scale,
        "primary": primary,
        "comparison": comparison,
    }


def _formula(
    ordinal: int,
    formula_type: str,
    expression_id: str,
    anchor_sheet: str,
    anchor_row: int,
    scope: str,
    bindings: Sequence[Mapping[str, Any]],
    dependencies: Sequence[str],
    expression_ast: list[Any],
    *,
    denominator_rule: Mapping[str, Any] | None = None,
    outputs: Sequence[Mapping[str, Any]],
    period_rule: str = "EXACT_FLOW_RANGE",
    disposition_rule: str = "EXACT_ZERO_VARIANCE_IN_BALANCE_ELSE_OUT_OF_BALANCE",
) -> dict[str, Any]:
    ordered_inputs = _with_ordinals(bindings)
    formula = {
        "formula_ordinal": ordinal,
        "formula_type": formula_type,
        "version": 1,
        "expression_id": expression_id,
        "anchor_sheet": anchor_sheet,
        "anchor_row": anchor_row,
        "scope": scope,
        "required": True,
        "expression_ast": expression_ast,
        "dependency_tuple": list(dependencies),
        "ordered_input_specs": ordered_inputs,
        "sign_rules": [
            [item["binding_ordinal"], item["sign_transform"]]
            for item in ordered_inputs
            if item["binding_kind"] == "SOURCE"
        ],
        "cardinality_rules": [
            [item["binding_ordinal"], item["cardinality"]] for item in ordered_inputs
        ],
        "period_rule": period_rule,
        "currency_rule": "EXACT_COMMON_ISO4217_NO_FX",
        "unit_rule": "EXACT_COMMON_SCALE_NO_CONVERSION",
        "denominator_rule": dict(denominator_rule) if denominator_rule else None,
        "output_specs": [dict(item) for item in outputs],
        "precision": 50,
        "rounding": "ROUND_HALF_UP",
        "permitted_assumptions": [],
        "disposition_rule": disposition_rule,
    }
    formula["formula_sha256"] = formula_sha256(formula)
    return formula


def formula_sha256(formula: Mapping[str, Any]) -> str:
    return sha256_canonical(
        [
            "fdd-formula-spec-v1",
            formula["formula_type"],
            formula["version"],
            formula["expression_id"],
            formula["expression_ast"],
            formula["dependency_tuple"],
            formula["ordered_input_specs"],
            formula["sign_rules"],
            formula["cardinality_rules"],
            formula["period_rule"],
            formula["currency_rule"],
            formula["unit_rule"],
            formula["denominator_rule"],
            formula["output_specs"],
            formula["precision"],
            formula["rounding"],
            formula["permitted_assumptions"],
        ]
    )


def expected_methodology_transition() -> dict[str, Any]:
    return {
        "schema_version": "fdd-methodology-transition-v1",
        "source_artifact_sha256": WORKBOOK_SHA256,
        "legacy_version_id": f"financial_dd:{WORKBOOK_SHA256}",
        "successor_version_id_preimage": [
            "financial_dd",
            WORKBOOK_SHA256,
            POLICY_ID,
            "$FDD_FORMULA_POLICY_V1_SHA256",
        ],
        "preserve_methodology_question_ids": True,
        "preserved_question_count": 79,
        "legacy_sheet_fallbacks": [
            {
                "source_sheet": "P&L",
                "calc_type": "GROSS_MARGIN",
                "required": False,
                "affected_question_count": 29,
            },
            {
                "source_sheet": "Cash Flow",
                "calc_type": "BURN_RATE",
                "required": False,
                "affected_question_count": 13,
            },
        ],
        "successor_required_map": [
            {
                "formula_ordinal": ordinal,
                "formula_type": formula_type,
                "source_sheet": sheet,
                "source_row": row,
                "required": True,
                "scope": scope,
            }
            for ordinal, formula_type, sheet, row, scope in FORMULA_ANCHORS
        ],
        "successor_required_question_count": 10,
        "successor_empty_requirement_count": 69,
        "strict_new_run_selection_only": True,
        "historical_run_selections_immutable": True,
    }


def expected_scope_exclusions() -> dict[str, Any]:
    return {
        "schema_version": "fdd-formula-scope-exclusions-v1",
        "rows": [
            {
                "scope_id": "fdd-debt-roll-forward-balance-scope-v1",
                "formula_type": "FDD_DEBT_ROLL_FORWARD",
                "included_balance_sources": [
                    {
                        "source_sheet": "Liabilities",
                        "source_row": 9,
                        "predicate": "fdd.interest_bearing_debt_balance",
                    }
                ],
                "excluded_balance_rows": [
                    {"source_sheet": "Liabilities", "source_row": 17},
                    {"source_sheet": "Liabilities", "source_row": 18},
                ],
                "included_movement_rows": [
                    {"source_sheet": "Cash Flow", "source_row": 11},
                    {"source_sheet": "Cash Flow", "source_row": 12},
                ],
                "excluded_rows_remain_required_answers": True,
                "unmatched_movement_disposition": "OUT_OF_BALANCE",
                "runtime_scope_expansion": "FORBIDDEN",
            },
            {
                "scope_id": "fdd-liability-equity-tie-out-scope-v1",
                "formula_type": "FDD_LIABILITY_TIE_OUT",
                "included_component_rows": [
                    {"source_sheet": "Liabilities", "source_row": row} for row in range(4, 19)
                ],
                "excluded_off_balance_sheet_rows": [
                    {"source_sheet": "Liabilities", "source_row": row} for row in range(19, 25)
                ],
                "excluded_rows_remain_required_answers": True,
                "manufactured_reported_total": "FORBIDDEN",
            },
        ],
    }


def expected_completeness_registry() -> dict[str, Any]:
    return {
        "schema_version": "fdd-documentary-completeness-requests-v1",
        "rows": [
            {
                "request_ordinal": ordinal,
                "predicate": predicate,
                "source_sheet": sheet,
                "source_row": row,
                "request_text": text,
                "admitted_result": "CITED_BOOLEAN_TRUE_OR_EVIDENCE_NOT_FOUND",
                "false_result": "FORBIDDEN",
                "uncited_true": "FORBIDDEN",
                "prose_only": "FORBIDDEN",
                "request_sha256": sha256_canonical(
                    [
                        "fdd-documentary-completeness-request-v1",
                        predicate,
                        sheet,
                        row,
                        text,
                    ]
                ),
            }
            for ordinal, (predicate, sheet, row, text) in enumerate(
                COMPLETENESS_REQUESTS,
                start=1,
            )
        ],
    }


def expected_direct_signed_registry() -> dict[str, Any]:
    def row(
        ordinal: int,
        predicate: str,
        positive: str,
        negative: str,
    ) -> dict[str, Any]:
        request = (
            f"Interpret a positive value as {positive}; interpret a negative value "
            f"as {negative}. Return the source's reported sign under this convention "
            "only; do not infer, balance, or silently negate it."
        )
        if predicate == "fdd.reported_delta_nwc":
            request = (
                "This predicate is an operating-NWC balance movement, not its cash effect. "
                f"Interpret a positive value as {positive}; interpret a negative value as "
                f"{negative}. The formula derives cash effect as -reported_delta_nwc; "
                "never return a cash-effect sign under this movement predicate. Return the "
                "source's reported sign under this convention only; do not infer, balance, "
                "or silently negate it."
            )
        return {
            "convention_ordinal": ordinal,
            "predicate": predicate,
            "positive_value": positive,
            "negative_value": negative,
            "sign_transform": "DIRECT_SIGNED",
            "extractor_request_requirement": request,
            "convention_sha256": sha256_canonical(
                [
                    "fdd-direct-signed-convention-v1",
                    predicate,
                    positive,
                    negative,
                    request,
                ]
            ),
        }

    return {
        "schema_version": "fdd-direct-signed-conventions-v1",
        "rows": [
            row(ordinal, predicate, positive, negative)
            for ordinal, (predicate, positive, negative) in enumerate(
                DIRECT_SIGNED_CONVENTIONS,
                start=1,
            )
        ],
        "reported_delta_nwc_cash_effect_rule": {
            "predicate": "fdd.reported_delta_nwc",
            "source_semantics": "MOVEMENT",
            "derived_cash_effect": "-reported_delta_nwc",
            "cash_effect_input_under_movement_predicate": "FORBIDDEN",
        },
    }


def expected_incomplete_schema() -> dict[str, Any]:
    return {
        "schema_version": "methodology-calculation-incomplete-v1",
        "code": "METHODOLOGY_CALCULATION_INCOMPLETE",
        "field_tuple": list(INCOMPLETE_TOP_FIELDS),
        "issue_field_tuple": list(INCOMPLETE_ISSUE_FIELDS),
        "issue_kind_enum": ["MISSING", "DUPLICATE", "AMBIGUOUS", "UNCONSUMED"],
        "missing_candidate_tuple": "EMPTY",
        "duplicate_ambiguous_candidates": "ALL_CONFLICTING_FACTUAL_OUTPUT_IDS",
        "unconsumed_candidates": "ALL_UNCONSUMED_GOVERNED_FACTUAL_OUTPUT_IDS",
        "issue_sort_tuple": [
            "formula_ordinal",
            "binding_ordinal",
            "issue_kind",
            "candidate_factual_output_ids",
        ],
        "resolution_mode": "COMPLETE_PASS_NEVER_FIRST_FAILURE",
        "private_fact_values": "FORBIDDEN",
        "claim_prose": "FORBIDDEN",
        "source_document_names": "FORBIDDEN",
    }


def expected_financial_output_schema() -> dict[str, Any]:
    return {
        "schema_version": "financial-diligence-row-set-v1",
        "top_level_field_tuple": list(FINANCIAL_TOP_FIELDS),
        "row_count": 79,
        "row_schema": {
            "schema_version": "fdd-dashboard-row-v1",
            "field_tuple": list(FINANCIAL_ROW_FIELDS),
            "result_record_field_tuple": list(FINANCIAL_RESULT_FIELDS),
            "answer_value_field_tuple": list(FINANCIAL_ANSWER_FIELDS),
            "evidence_link_field_tuple": [
                "evidence_link_id",
                "evidence_id",
                "output_object_id",
                "output_identity_sha256",
                "output_value_sha256",
            ],
            "sanad_link_field_tuple": [
                "sanad_link_id",
                "sanad_id",
                "grade",
                "output_object_id",
                "output_identity_sha256",
                "output_value_sha256",
            ],
            "task_gap_field_tuple": list(FINANCIAL_TASK_GAP_FIELDS),
            "root_link_field_tuple": list(FINANCIAL_ROOT_LINK_FIELDS),
            "calculation_package": (
                "FddDeterministicCalculationV1_PLUS_CALC_SANAD_AND_NATIVE_HASHES"
            ),
        },
        "reconciliation_field_tuple": list(FINANCIAL_RECONCILIATION_FIELDS),
        "row_payload_preimage": [
            "financial-diligence-row-v1",
            "$identity_and_trace_tuple",
            "$disposition",
            "$disposition_reason_code",
            "$result_records",
            "$answer_values",
            "$task_gaps",
            "$calculations",
            "$root_projection_sha256",
        ],
        "identity_and_trace_tuple": [
            "row_ordinal",
            "sheet_ordinal",
            "source_sheet",
            "source_row_number",
            "source_anchor_cell",
            "methodology_version_record_id",
            "methodology_version_id",
            "methodology_question_record_id",
            "methodology_question_id",
            "question_payload_hash",
            "coverage_record_id",
            "dashboard_row_id",
            "dashboard_output_object_id",
        ],
        "row_set_preimage_tuple": [
            "ordinal",
            "dashboard_row_id",
            "row_payload_sha256",
            "root_projection_sha256",
        ],
        "calculation_set_preimage_tuple": [
            "policy_ordinal",
            "coordinate",
            "calc_id",
            "reproducibility_sha256",
        ],
        "sort_contract": {
            "rows": [
                "sheet_ordinal",
                "source_row_number",
                "methodology_question_id",
            ],
            "results": ["task_record_id", "result_id"],
            "answer_values": ["output_ordinal", "factual_output_record_id"],
            "task_gaps": ["task_record_id", "result_id"],
            "root_links": [
                "root_kind",
                "root_ordinal",
                "output_object_id",
                "root_link_id",
            ],
            "calculations": ["policy_ordinal", "coordinate", "calc_id"],
            "inputs": ["binding_ordinal"],
        },
        "unknown_fields": "FORBIDDEN",
        "embedded_self_hash": "FORBIDDEN",
    }


def expected_activation_assertions() -> dict[str, Any]:
    return {
        "schema_version": "fdd-formula-policy-activation-v1",
        "policy_status": "INACTIVE",
        "append_frozen": True,
        "runtime_selectable": False,
        "migration_consumable": False,
        "formula_implementation_authorized": False,
        "repository_implementation_authorized": False,
        "worker_implementation_authorized": False,
        "dashboard_implementation_authorized": False,
        "deliverable_implementation_authorized": False,
        "second_artifact_fixture_review_required": True,
        "successor_pipeline_contract": "full-strict-v22",
        "successor_pipeline_contract_status": "INACTIVE_PENDING_REVIEW",
        "legacy_pipeline_contract": "full-strict-v21",
        "legacy_pipeline_contract_immutable": True,
        "migration": "0039",
        "migration_source_authorized": False,
        "live_calls_authorized": False,
    }


def expected_formula_registry() -> dict[str, Any]:
    return {
        "schema_version": "fdd-formula-registry-v1",
        "formula_hash_preimage_field_tuple": [
            "formula_type",
            "version",
            "expression_id",
            "expression_ast",
            "dependency_tuple",
            "ordered_input_specs",
            "sign_rules",
            "cardinality_rules",
            "period_rule",
            "currency_rule",
            "unit_rule",
            "denominator_rule",
            "output_specs",
            "precision",
            "rounding",
            "permitted_assumptions",
        ],
        "rows": expected_formula_rows(),
    }


def recompute_component_hashes(policy: Mapping[str, Any]) -> dict[str, str]:
    return {
        section: sha256_canonical(["fdd-formula-policy-v1-component-v1", section, policy[section]])
        for section in HASHED_COMPONENTS
    }


def recompute_policy_hash(policy: Mapping[str, Any]) -> str:
    component_hashes = recompute_component_hashes(policy)
    return sha256_canonical(
        [
            "fdd-formula-policy-v1",
            [[section, component_hashes[section]] for section in HASHED_COMPONENTS],
        ]
    )


def expected_digest_registry(policy: Mapping[str, Any]) -> dict[str, Any]:
    component_hashes = recompute_component_hashes(policy)
    return {
        "schema_version": "fdd-formula-policy-digests-v1",
        "algorithm": "SHA-256",
        "component_order": list(HASHED_COMPONENTS),
        "component_hashes": component_hashes,
        "policy_hash_preimage": [
            "fdd-formula-policy-v1",
            [[section, component_hashes[section]] for section in HASHED_COMPONENTS],
        ],
        "fdd_formula_policy_v1_sha256": recompute_policy_hash(policy),
    }


def expected_policy() -> dict[str, Any]:
    policy = {
        "metadata": expected_metadata(),
        "canonicalization_contract": expected_canonicalization_contract(),
        "workbook_authority": expected_workbook_authority(),
        "workbook_row_registry": {
            "schema_version": "fdd-workbook-row-registry-v1",
            "rows": read_operator_rows(),
        },
        "methodology_transition": expected_methodology_transition(),
        "formula_registry": expected_formula_registry(),
        "debt_scope_exclusions": expected_scope_exclusions(),
        "completeness_request_registry": expected_completeness_registry(),
        "direct_signed_convention_registry": expected_direct_signed_registry(),
        "incomplete_binding_failure_schema": expected_incomplete_schema(),
        "financial_output_schema": expected_financial_output_schema(),
        "digest_recipe_registry": {},
        "activation_assertions": expected_activation_assertions(),
    }
    policy["digest_recipe_registry"] = expected_digest_registry(policy)
    return policy


def validate_policy(policy: Mapping[str, Any]) -> None:
    validate_top_level(policy)
    _safe_require(
        policy["canonicalization_contract"] == expected_canonicalization_contract(),
        "TASK8_FDD_CANONICALIZATION_INVALID",
    )
    _safe_require(
        policy["workbook_authority"] == expected_workbook_authority(),
        "TASK8_FDD_WORKBOOK_AUTHORITY_INVALID",
    )
    expected_rows = read_operator_rows()
    _safe_require(
        policy["workbook_row_registry"]
        == {
            "schema_version": "fdd-workbook-row-registry-v1",
            "rows": expected_rows,
        },
        "TASK8_FDD_WORKBOOK_ROW_REGISTRY_INVALID",
    )
    _safe_require(
        policy["methodology_transition"] == expected_methodology_transition(),
        "TASK8_FDD_METHODOLOGY_TRANSITION_INVALID",
    )
    _safe_require(
        policy["formula_registry"] == expected_formula_registry(),
        "TASK8_FDD_FORMULA_REGISTRY_INVALID",
    )
    _safe_require(
        policy["debt_scope_exclusions"] == expected_scope_exclusions(),
        "TASK8_FDD_SCOPE_EXCLUSIONS_INVALID",
    )
    _safe_require(
        policy["completeness_request_registry"] == expected_completeness_registry(),
        "TASK8_FDD_COMPLETENESS_REQUESTS_INVALID",
    )
    _safe_require(
        policy["direct_signed_convention_registry"] == expected_direct_signed_registry(),
        "TASK8_FDD_DIRECT_SIGNED_CONVENTIONS_INVALID",
    )
    _safe_require(
        policy["incomplete_binding_failure_schema"] == expected_incomplete_schema(),
        "TASK8_FDD_INCOMPLETE_SCHEMA_INVALID",
    )
    _safe_require(
        policy["financial_output_schema"] == expected_financial_output_schema(),
        "TASK8_FDD_FINANCIAL_OUTPUT_SCHEMA_INVALID",
    )
    _safe_require(
        policy["activation_assertions"] == expected_activation_assertions(),
        "TASK8_FDD_ACTIVATION_ASSERTIONS_INVALID",
    )
    _safe_require(
        policy["digest_recipe_registry"] == expected_digest_registry(policy),
        "TASK8_FDD_DIGEST_REGISTRY_INVALID",
    )
    validate_derived_counts(policy)


def validate_derived_counts(policy: Mapping[str, Any]) -> None:
    formula_rows = policy["formula_registry"]["rows"]
    source_bindings = [
        binding
        for formula in formula_rows
        for binding in formula["ordered_input_specs"]
        if binding["binding_kind"] == "SOURCE"
    ]
    parent_bindings = [
        binding
        for formula in formula_rows
        for binding in formula["ordered_input_specs"]
        if binding["binding_kind"] == "PARENT_CALCULATION"
    ]
    distinct_source_mappings = {
        (binding["source_sheet"], binding["source_row"], binding["predicate"])
        for binding in source_bindings
    }
    direct_signed = [
        binding for binding in source_bindings if binding["sign_transform"] == "DIRECT_SIGNED"
    ]
    completeness = [
        binding for binding in source_bindings if binding["value_schema"] == "FddBooleanValueV1"
    ]
    _safe_require(len(formula_rows) == 10, "TASK8_FDD_FORMULA_COUNT_INVALID")
    _safe_require(len(source_bindings) == 107, "TASK8_FDD_SOURCE_BINDING_COUNT_INVALID")
    _safe_require(
        len(distinct_source_mappings) == 101,
        "TASK8_FDD_DISTINCT_SOURCE_MAPPING_COUNT_INVALID",
    )
    _safe_require(len(parent_bindings) == 9, "TASK8_FDD_PARENT_BINDING_COUNT_INVALID")
    _safe_require(
        len(direct_signed) == 19,
        "TASK8_FDD_DIRECT_SIGNED_BINDING_COUNT_INVALID",
    )
    _safe_require(
        len(completeness) == 2,
        "TASK8_FDD_COMPLETENESS_BINDING_COUNT_INVALID",
    )
    for formula in formula_rows:
        _safe_require(
            [item["binding_ordinal"] for item in formula["ordered_input_specs"]]
            == list(range(1, len(formula["ordered_input_specs"]) + 1)),
            "TASK8_FDD_BINDING_ORDINAL_INVALID",
        )
        _safe_require(
            formula["formula_sha256"] == formula_sha256(formula),
            "TASK8_FDD_FORMULA_HASH_INVALID",
        )


def render_review_map(policy: Mapping[str, Any]) -> bytes:
    """Render the deterministic reviewer projection without logging row prose."""

    workbook_rows = {
        (row["source_sheet"], row["source_row"]): row
        for row in policy["workbook_row_registry"]["rows"]
    }
    lines = [
        "# Slice100 Task 8 FDD Formula Policy v1 — reviewer mapping",
        "",
        "<!-- Deterministic projection of contracts/fdd_formula_policy_v1.json. -->",
        "",
        f"- Policy ID: `{POLICY_ID}`",
        (f"- Policy SHA-256: `{policy['digest_recipe_registry']['fdd_formula_policy_v1_sha256']}`"),
        f"- Workbook SHA-256: `{WORKBOOK_SHA256}`",
        "- Formula anchors: `10`",
        "- Expanded source bindings: `107`",
        "",
    ]
    for formula in policy["formula_registry"]["rows"]:
        anchor = workbook_rows[(formula["anchor_sheet"], formula["anchor_row"])]
        lines.extend(
            [
                (f"## {formula['formula_ordinal']}. {_markdown_cell(formula['formula_type'])}"),
                "",
                "| Anchor sheet | Row | Scope | Term | Nature | Line item | Question | "
                "Anchor cell | Methodology question ID | Source-row SHA-256 |",
                "|---|---:|---|---|---|---|---|---|---|---|",
                (
                    f"| {_markdown_cell(anchor['source_sheet'])} | {anchor['source_row']} | "
                    f"{_markdown_cell(formula['scope'])} | {_markdown_cell(anchor['term'])} | "
                    f"{_markdown_cell(anchor['nature'])} | "
                    f"{_markdown_cell(anchor['line_item'])} | "
                    f"{_markdown_cell(anchor['question_text'])} | "
                    f"{_markdown_cell(anchor['source_anchor_cell'])} | "
                    f"`{anchor['methodology_question_id']}` | "
                    f"`{anchor['source_row_sha256']}` |"
                ),
                "",
                "| Binding | Role | Predicate | Sign | Cardinality | Sheet | Row | Term | "
                "Nature | Line item | Question | Anchor cell | Methodology question ID | "
                "Mapping SHA-256 |",
                "|---:|---|---|---|---|---|---:|---|---|---|---|---|---|---|",
            ]
        )
        for binding in formula["ordered_input_specs"]:
            if binding["binding_kind"] != "SOURCE":
                continue
            row = workbook_rows[(binding["source_sheet"], binding["source_row"])]
            mapping_hash = sha256_canonical(
                [
                    "fdd-formula-source-mapping-v1",
                    formula["formula_ordinal"],
                    binding,
                    row["methodology_question_id"],
                    row["source_row_sha256"],
                ]
            )
            lines.append(
                f"| {binding['binding_ordinal']} | {_markdown_cell(binding['role'])} | "
                f"`{binding['predicate']}` | `{binding['sign_transform']}` | "
                f"`{binding['cardinality']}` | {_markdown_cell(row['source_sheet'])} | "
                f"{row['source_row']} | {_markdown_cell(row['term'])} | "
                f"{_markdown_cell(row['nature'])} | {_markdown_cell(row['line_item'])} | "
                f"{_markdown_cell(row['question_text'])} | "
                f"{_markdown_cell(row['source_anchor_cell'])} | "
                f"`{row['methodology_question_id']}` | `{mapping_hash}` |"
            )
        lines.append("")
    return ("\n".join(lines).rstrip() + "\n").encode("utf-8")


def _markdown_cell(value: object) -> str:
    if value is None:
        return "—"
    return (
        str(value)
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\r\n", "<br>")
        .replace("\n", "<br>")
        .replace("\r", "<br>")
    )
